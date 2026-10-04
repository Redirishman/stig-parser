"""End-to-end through parse_stage: the defects in spec §1 must stay fixed."""
import ast
import json
import logging
import tempfile
import zipfile
from pathlib import Path

import pytest

from app.core.pipeline import PipelineError, parse_stage

FIX = Path(__file__).parent / "fixtures"


def test_manual_stig_fills_text_for_long_form_scan_ids(tmp_path):
    # Defect 2: a real Manual STIG matched zero rules.
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [FIX / "manual_stig_server2022.xml"], tmp_path)
    f = {x.vuln_id or x.rule_id: x for x in result.findings}
    drifted = f["V-254239"]
    assert drifted.check_text and drifted.fix_text and drifted.severity == "CAT I"
    assert drifted.stig_id == "WN22-00-000010"
    assert drifted.text_source == "Check and fix: manual_stig_server2022.xml V2R8, revision differs from scan"
    same = f["V-254241"]
    assert same.check_text and same.text_source == "Check and fix: manual_stig_server2022.xml V2R8"
    assert all(x.stig_title == "Microsoft Windows Server 2022 Security Technical Implementation Guide"
               for x in result.findings)


def test_unmatched_rules_reach_the_operator(tmp_path):
    # Defect 3: the mismatch was logged, never returned.
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [FIX / "manual_stig_server2022.xml"], tmp_path)
    assert any("not found in any supplied reference" in w and "SV-254243" in w for w in result.warnings)
    assert any("different release" in w for w in result.warnings)


def test_reference_upload_does_not_wipe_scc_data(tmp_path):
    # Defect 4: supplying any benchmark dropped the embedded one.
    alone = parse_stage([FIX / "scc_embedded_results.xml"], [], tmp_path / "a")
    with_ref = parse_stage([FIX / "scc_embedded_results.xml"], [FIX / "manual_stig_win11.xml"], tmp_path / "b")
    for a, b in zip(alone.findings, with_ref.findings):
        assert (b.fix_text, b.severity, b.vuln_id, b.stig_title) == (a.fix_text, a.severity, a.vuln_id, a.stig_title)
        assert a.fix_text and a.severity and a.vuln_id
    filled = [f for f in with_ref.findings if f.check_text]
    assert {f.vuln_id for f in filled} == {"V-253284", "V-253285"}


def test_plain_scc_run_says_why_check_text_is_blank(tmp_path):
    result = parse_stage([FIX / "scc_embedded_results.xml"], [], tmp_path)
    assert all(not f.check_text and f.fix_text for f in result.findings)
    assert all(f.text_source == "Check: no reference supplied | Fix: scanner" for f in result.findings)
    assert [w for w in result.warnings if "Check text is blank for 3 finding(s)" in w]


def test_wrong_slot_still_works(tmp_path):
    # The SPA sends one flat list: a Manual STIG among the results must act as a reference.
    result = parse_stage([FIX / "scc_embedded_results.xml", FIX / "manual_stig_win11.xml"], [], tmp_path)
    assert result.source_file_count == 1
    assert any(f.check_text for f in result.findings)
    assert not any("0 rule results" in w for w in result.warnings)


def test_cklb_reference_fills_xccdf_findings(tmp_path):
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [FIX / "evaluate_stig_checklist.cklb"], tmp_path)
    assert result.source_file_count == 1                      # the checklist is a reference, not a second scan
    assert all(f.check_text and f.fix_text for f in result.findings)


def test_scap_datastream_reference_supplies_fix_text_and_title(tmp_path):
    result = parse_stage([FIX / "scc_embedded_results.xml"], [FIX / "scap_datastream_win11.xml"], tmp_path)
    assert all(f.stig_title for f in result.findings)
    assert any("Check text is blank" in w or "not found in any supplied reference" in w for w in result.warnings)


def test_enrichment_report_and_coverage_are_returned(tmp_path):
    result = parse_stage([FIX / "scc_embedded_results.xml"], [FIX / "manual_stig_win11.xml"], tmp_path)
    assert [s.file_name for s in result.enrichment.standalone_sources] == ["manual_stig_win11.xml"]
    assert {(f.server, f.stig_title) for f in result.findings} <= result.coverage


def test_references_only_is_an_error_that_says_so(tmp_path):
    with pytest.raises(PipelineError) as exc:
        parse_stage([FIX / "manual_stig_win11.xml"], [], tmp_path)
    assert "no rule results" in str(exc.value).lower() and "reference" in str(exc.value).lower()


# --- each scan is read against the benchmark in its own results file -------------------

def test_text_from_an_operator_reference_is_never_labelled_scanner(tmp_path):
    # The results file carries no benchmark: every value on the row comes from the reference.
    result = parse_stage([FIX / "scc_results.xml"], [FIX / "sample_benchmark.xml"], tmp_path)
    assert len(result.findings) == 5
    for f in result.findings:
        assert f.check_text and f.fix_text and f.vuln_id and f.severity.startswith("CAT ")
        assert "scanner" not in f.text_source
        assert f.text_source == "Check and fix: sample_benchmark.xml V1R4"
        assert f.scan_release == ""                 # the reference's release is not what was scanned
    assert not any("no matching STIG benchmark" in w for w in result.warnings)


def test_the_scanned_release_stays_the_embedded_one_when_a_newer_reference_is_supplied(tmp_path):
    result = parse_stage([FIX / "scc_embedded_results.xml"], [FIX / "manual_stig_win11.xml"], tmp_path)
    assert {f.scan_release for f in result.findings} == {"V2R8"}
    by_vuln = {f.vuln_id: f for f in result.findings}
    assert by_vuln["V-253284"].text_source == "Check: manual_stig_win11.xml V2R9 | Fix: scanner"
    assert by_vuln["V-253285"].text_source == "Check: manual_stig_win11.xml V2R9, scanned V2R8 | Fix: scanner"
    assert by_vuln["V-253286"].text_source == "Check: not in supplied references | Fix: scanner"
    assert all(f.stig_title == "Microsoft Windows 11 STIG SCAP Benchmark" for f in result.findings)


def test_a_rule_found_in_a_reference_without_check_text_says_so(tmp_path):
    # The SCAP datastream holds SV-253284 but carries no check text: the rule was found.
    result = parse_stage([FIX / "evaluate_stig_results.xml", FIX / "scc_embedded_results.xml"],
                         [FIX / "scap_datastream_win11.xml"], tmp_path)
    by_vuln = {f.vuln_id: f for f in result.findings}
    assert by_vuln["V-253284"].text_source == \
        "Check: in scap_datastream_win11.xml V2R8, which has no check text | Fix: scanner"
    assert "not found in any supplied reference" not in " ".join(
        w for w in result.warnings if "SV-253284" in w)


def _cell_tally(findings) -> dict[str, dict[str, int]]:
    """Per title, the counts the Text Source cells imply."""
    tally: dict[str, dict[str, int]] = {}
    for f in findings:
        t = tally.setdefault(f.stig_title or "(no STIG title)", {"unmatched": 0, "ambiguous": 0, "product_refused": 0})
        t["unmatched"] += "not in supplied references" in f.text_source
        t["ambiguous"] += "matches several" in f.text_source
        t["product_refused"] += "under a different STIG title" in f.text_source
    return tally


@pytest.mark.parametrize("references", [
    ["scap_datastream_win11.xml"],
    ["scap_datastream_win11.xml", "manual_stig_server2022.xml"],
    ["manual_stig_win11.xml"],
])
def test_the_counts_and_the_cells_agree(tmp_path, references):
    result = parse_stage([FIX / "evaluate_stig_results.xml", FIX / "scc_embedded_results.xml"],
                         [FIX / name for name in references], tmp_path)
    counts = {t: {"unmatched": c.unmatched, "ambiguous": c.ambiguous, "product_refused": c.product_refused}
              for t, c in result.enrichment.stigs.items()}
    assert _cell_tally(result.findings) == counts


def _scc_copy(directory: Path, host: str, version: str, name: str = "scan.xml") -> Path:
    """The SCC fixture as another host's scan of another release, under the SAME file name by default."""
    text = (FIX / "scc_embedded_results.xml").read_text(encoding="utf-8")
    assert "<cdf:version>002.008</cdf:version>" in text and "<cdf:target>WKSTN-01</cdf:target>" in text
    text = text.replace("<cdf:version>002.008</cdf:version>", f"<cdf:version>{version}</cdf:version>")
    text = text.replace("<cdf:target>WKSTN-01</cdf:target>", f"<cdf:target>{host}</cdf:target>")
    assert ">Set the registry value" in text
    text = text.replace(">Set the registry value", f">Fix text of {version}: set the registry value")
    directory.mkdir()
    path = directory / name
    path.write_text(text, encoding="utf-8")
    return path


def test_two_results_files_with_one_name_each_keep_their_own_benchmark(tmp_path):
    # Two uploads (two ZIPs, two folders) can hold the same file name.
    older = _scc_copy(tmp_path / "a", "HOST-A", "002.008")
    newer = _scc_copy(tmp_path / "b", "HOST-B", "002.009")
    result = parse_stage([older, newer], [], tmp_path / "x")
    assert {(f.server, f.scan_release) for f in result.findings} == {("HOST-A", "V2R8"), ("HOST-B", "V2R9")}
    sehop = {f.server: f.fix_text for f in result.findings if f.vuln_id == "V-253284"}
    assert sehop["HOST-A"].startswith("Fix text of 002.008: ") and sehop["HOST-B"].startswith("Fix text of 002.009: ")
    assert len(result.findings) == 6 and result.source_file_count == 2
    assert all(f.fix_text and f.text_source.endswith("| Fix: scanner") for f in result.findings)


def test_identical_embedded_benchmarks_are_one_source_and_every_scan_still_has_its_own(tmp_path):
    first = _scc_copy(tmp_path / "a", "HOST-A", "002.008")
    second = _scc_copy(tmp_path / "b", "HOST-B", "002.008")
    result = parse_stage([first, second], [], tmp_path / "x")
    assert len(result.enrichment.sources) == 1                      # the library keeps one copy
    assert {(f.server, f.scan_release) for f in result.findings} == {("HOST-A", "V2R8"), ("HOST-B", "V2R8")}
    assert all(f.fix_text and f.vuln_id and f.severity for f in result.findings)
    assert [w for w in result.warnings if "Check text is blank for 6 finding(s) in 1 STIG(s)" in w]


def test_identical_embedded_benchmarks_under_different_names_each_serve_their_own_scan(tmp_path):
    # The library keeps one copy and answers the second file with the FIRST file's source.
    # Which source it returned must not decide whether a scan gets its own rule data.
    first = _scc_copy(tmp_path / "a", "HOST-A", "002.008", name="host-a.xml")
    second = _scc_copy(tmp_path / "b", "HOST-B", "002.008", name="host-b.xml")
    result = parse_stage([first, second], [], tmp_path / "x")
    assert [s.file_name for s in result.enrichment.sources] == ["host-a.xml"]
    assert sorted(f.server for f in result.findings) == ["HOST-A"] * 3 + ["HOST-B"] * 3
    for f in result.findings:
        assert f.fix_text and f.vuln_id and f.severity.startswith("CAT ") and f.stig_id
        assert (f.scan_release, f.stig_title) == ("V2R8", "Microsoft Windows 11 STIG SCAP Benchmark")
        assert f.text_source == "Check: no reference supplied | Fix: scanner"
    assert not any("no matching STIG benchmark" in w for w in result.warnings)


def test_a_scan_with_no_benchmark_anywhere_is_reported_by_file(tmp_path):
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [], tmp_path)
    assert all(f.text_source == "Check and fix: no reference supplied" for f in result.findings)
    assert all(f.severity == "" for f in result.findings)
    assert (
        "evaluate_stig_results.xml: no matching STIG benchmark — 3 finding(s) have no STIG title, "
        "severity, or check/fix text. Add the STIG as a reference."
    ) in result.warnings


def test_a_results_checklist_also_fills_another_scanners_findings(tmp_path):
    result = parse_stage([FIX / "evaluate_stig_results.xml", FIX / "evaluate_stig_checklist.cklb"], [], tmp_path)
    assert result.source_file_count == 2
    from_xccdf = [f for f in result.findings if f.rule_id.startswith("xccdf_")]
    assert len(from_xccdf) == 3
    assert all(f.check_text and f.fix_text and f.stig_title == "Microsoft Windows Server 2022 STIG" for f in from_xccdf)
    assert all("evaluate_stig_checklist.cklb V1R4" in f.text_source and "scanner" not in f.text_source
               for f in from_xccdf)
    from_checklist = [f for f in result.findings if not f.rule_id.startswith("xccdf_")]
    assert from_checklist and all(f.text_source == "Check and fix: scanner" for f in from_checklist)
    assert not any("no matching STIG benchmark" in w for w in result.warnings)


# --- what the library and the loaders could not use reaches the operator ----------------

def _crowded_manual_stig(path: Path) -> Path:
    rules = "".join(
        f'<Rule id="SV-{n}r1_rule" severity="medium"><version>X-{n}</version><title>r</title>'
        f"<fixtext>fix</fixtext><check><check-content>check</check-content></check></Rule>"
        for n in range(1, 41))
    path.write_text(
        '<Benchmark xmlns="http://checklists.nist.gov/xccdf/1.1" id="X_STIG"><title>X STIG</title>'
        f'<version>1</version><Group id="V-1">{rules}</Group></Benchmark>', encoding="utf-8")
    return path


def test_library_warnings_reach_the_operator(tmp_path):
    reference = _crowded_manual_stig(tmp_path / "crowded.xml")
    result = parse_stage([FIX / "scc_embedded_results.xml"], [reference], tmp_path / "x")
    assert "crowded.xml: more than 32 rules share one V-ID — extra rules ignored for lookup" in result.warnings


def _checklist(path: Path, stigs) -> Path:
    path.write_text(json.dumps({"stigs": stigs}), encoding="utf-8")
    return path


def test_an_unreadable_reference_checklist_is_named(tmp_path):
    bad = tmp_path / "bad.cklb"
    bad.write_text("{not json", encoding="utf-8")
    other = _checklist(tmp_path / "other.cklb", [])
    other.write_text('{"hello": 1}', encoding="utf-8")
    result = parse_stage([FIX / "scc_embedded_results.xml"], [bad, other], tmp_path / "x")
    assert ("Could not parse reference checklist: bad.cklb — not valid JSON: Expecting property name enclosed "
            "in double quotes: line 1 column 2 (char 1)") in result.warnings
    assert "Could not parse reference checklist: other.cklb — not a CKLB checklist (no 'stigs' array)" in result.warnings
    assert not any("0 rules" in w for w in result.warnings)


def test_an_empty_reference_checklist_is_ignored_and_said_to_be(tmp_path):
    empty = _checklist(tmp_path / "empty.cklb", [{"stig_id": "X", "display_name": "X STIG", "version": 1, "rules": []}])
    result = parse_stage([FIX / "scc_embedded_results.xml"], [empty], tmp_path / "x")
    assert "empty.cklb: reference checklist contains 0 rules — ignored" in result.warnings
    assert not any("Could not parse" in w for w in result.warnings)
    # Ignored means it does not count as a supplied reference.
    assert result.enrichment.standalone_sources == []
    assert all(f.text_source == "Check: no reference supplied | Fix: scanner" for f in result.findings)


def test_non_text_values_in_a_reference_checklist_are_counted_for_the_operator(tmp_path):
    rules = [{"group_id": "V-253284", "rule_id_src": "SV-253284r958928_rule", "rule_version": "WN11-00-000150",
              "check_content": {"nested": "object"}, "fix_text": ["a", "list"], "severity": 3}]
    odd = _checklist(tmp_path / "odd.cklb", [{"stig_id": "Microsoft_Windows_11_STIG", "version": 2,
                                              "display_name": "Microsoft Windows 11 STIG", "rules": rules}])
    result = parse_stage([FIX / "scc_embedded_results.xml"], [odd], tmp_path / "x")
    assert "odd.cklb: 3 non-text value(s) in the checklist were ignored" in result.warnings
    assert "odd.cklb: 0 of 1 rules carry check text" in result.warnings
    assert all("{" not in f.check_text and "nested" not in f.check_text for f in result.findings)


def test_a_reference_with_no_rules_is_ignored_and_said_to_be(tmp_path):
    hollow = tmp_path / "hollow.xml"
    hollow.write_text('<Benchmark xmlns="http://checklists.nist.gov/xccdf/1.1" id="X"><title>X STIG</title></Benchmark>',
                      encoding="utf-8")
    result = parse_stage([FIX / "scc_embedded_results.xml"], [hollow], tmp_path / "x")
    assert "hollow.xml: reference contains 0 rules — ignored" in result.warnings
    assert result.enrichment.standalone_sources == []


def test_a_reference_that_is_not_a_benchmark_is_named(tmp_path):
    unknown = tmp_path / "unknown.xml"
    unknown.write_text("<report><row/></report>", encoding="utf-8")
    result = parse_stage([FIX / "scc_embedded_results.xml"], [unknown], tmp_path / "x")
    assert "Could not parse benchmark: unknown.xml — no Benchmark element" in result.warnings


CKL = "STIG Viewer .ckl checklists are not supported — save it as .cklb in STIG Viewer 3 and upload that"


@pytest.mark.parametrize("slot", ["results", "reference"])
def test_a_loose_legacy_checklist_is_named_as_unsupported(tmp_path, slot):
    legacy = FIX / "legacy_checklist.ckl.xml"
    scan = FIX / "scc_embedded_results.xml"
    result = parse_stage([scan, legacy], [], tmp_path) if slot == "results" else parse_stage([scan], [legacy], tmp_path)
    assert result.source_file_count == 1 and len(result.findings) == 3
    assert [w for w in result.warnings if "legacy_checklist" in w] == [f"legacy_checklist.ckl.xml: {CKL}"]
    assert not any("Could not parse" in w for w in result.warnings)


def test_legacy_checklists_zipped_with_a_scan_are_named_as_unsupported(tmp_path):
    session = _zip_of(tmp_path / "session.zip", {
        "WKSTN-01_results.xml": FIX / "scc_embedded_results.xml",
        "legacy_checklist.ckl.xml": FIX / "legacy_checklist.ckl.xml",
        "OLD-HOST.ckl": FIX / "legacy_checklist.ckl.xml",
    })
    result = parse_stage([session], [], tmp_path / "x")
    assert result.source_file_count == 1 and len(result.findings) == 3
    assert f"legacy_checklist.ckl.xml: {CKL}" in result.warnings and f"OLD-HOST.ckl: {CKL}" in result.warnings


def test_only_legacy_checklists_is_an_error_that_carries_the_reason(tmp_path):
    with pytest.raises(PipelineError) as exc:
        parse_stage([FIX / "legacy_checklist.ckl.xml"], [], tmp_path)
    assert exc.value.warnings == [f"legacy_checklist.ckl.xml: {CKL}"]


def test_an_error_still_carries_the_reference_warnings(tmp_path):
    bad = tmp_path / "bad.cklb"
    bad.write_text("{not json", encoding="utf-8")
    broken = tmp_path / "broken.xml"
    broken.write_text("<TestResult><unclosed>", encoding="utf-8")
    with pytest.raises(PipelineError) as exc:
        parse_stage([broken], [bad, _crowded_manual_stig(tmp_path / "crowded.xml")], tmp_path / "x")
    assert "no valid results files" in str(exc.value).lower()
    assert ("Could not parse reference checklist: bad.cklb — not valid JSON: Expecting property name enclosed "
            "in double quotes: line 1 column 2 (char 1)") in exc.value.warnings
    assert any("more than 32 rules share one V-ID" in w for w in exc.value.warnings)


# --- a reference checklist names the STIG for a scan that carries no benchmark -----------

SERVER_2022 = "Microsoft Windows Server 2022 STIG"


def test_a_checklist_reference_names_the_scan_and_the_coverage_pair_carries_the_title(tmp_path):
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [FIX / "evaluate_stig_checklist.cklb"], tmp_path)
    assert {f.stig_title for f in result.findings} == {SERVER_2022}
    assert result.coverage == {("WIN-SERVER-04", SERVER_2022)}        # and no ("WIN-SERVER-04", "") pair
    assert not any("no matching STIG benchmark" in w for w in result.warnings)


def _one_rule_checklist(path: Path, stig_id="MS_Windows_Server_2022_STIG", title=SERVER_2022) -> Path:
    """A checklist of the Server 2022 STIG that holds one of the three rules the scan reports."""
    rules = [{"group_id": "V-254239", "rule_id_src": "SV-254239r958472_rule", "rule_version": "WN22-00-000010",
              "severity": "high", "check_content": "check", "fix_text": "fix"}]
    return _checklist(path, [{"stig_id": stig_id, "display_name": title, "version": 1,
                              "release_info": "Release: 4", "rules": rules}])


def test_a_finding_the_checklist_lacks_is_still_titled_and_reported_under_its_stig(tmp_path):
    reference = _one_rule_checklist(tmp_path / "partial.cklb")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [reference], tmp_path / "x")
    assert len(result.findings) == 3 and {f.stig_title for f in result.findings} == {SERVER_2022}
    filled = [f for f in result.findings if f.check_text]
    assert [f.vuln_id for f in filled] == ["V-254239"]
    assert all(f.text_source == "Check and fix: not in supplied references" for f in result.findings if not f.check_text)
    assert any(w.startswith(f"{SERVER_2022}: 2 finding(s) across 2 rule(s) not found in any supplied reference")
               for w in result.warnings), result.warnings
    assert not any("(no STIG title)" in w or "no matching STIG benchmark" in w for w in result.warnings)
    assert result.coverage == {("WIN-SERVER-04", SERVER_2022)}
    assert {f.scan_release for f in result.findings} == {""}          # the checklist's release was not scanned


def test_an_all_pass_scan_is_covered_under_the_title_a_checklist_gives(tmp_path):
    text = (FIX / "evaluate_stig_results.xml").read_text(encoding="utf-8")
    for status in ("fail", "notchecked", "unknown"):
        assert f"<result>{status}</result>" in text
        text = text.replace(f"<result>{status}</result>", "<result>pass</result>")
    clean = tmp_path / "clean.xml"
    clean.write_text(text, encoding="utf-8")
    result = parse_stage([clean], [FIX / "evaluate_stig_checklist.cklb"], tmp_path / "x", allow_empty=True)
    assert result.findings == []
    assert result.coverage == {("WIN-SERVER-04", SERVER_2022)}


def test_the_scans_own_benchmark_still_names_it_when_a_checklist_is_supplied(tmp_path):
    reference = _checklist(tmp_path / "win11.cklb", [{
        "stig_id": "Microsoft_Windows_11_STIG", "display_name": "Title from the checklist", "version": 2,
        "rules": [{"group_id": "V-253284", "rule_id_src": "SV-253284r958928_rule", "check_content": "check"}]}])
    result = parse_stage([FIX / "scc_embedded_results.xml"], [reference], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {"Microsoft Windows 11 STIG SCAP Benchmark"}
    assert {f.scan_release for f in result.findings} == {"V2R8"}


def test_a_checklist_with_no_display_name_names_the_scan_by_the_title_the_loader_gives_it(tmp_path):
    # The loader falls back to the STIG ID for the title, and the rules it fills carry the same one.
    reference = _one_rule_checklist(tmp_path / "untitled.cklb", title="")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [reference], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {"MS_Windows_Server_2022_STIG"}
    assert result.coverage == {("WIN-SERVER-04", "MS_Windows_Server_2022_STIG")}


@pytest.mark.parametrize("stig_id,title", [("", SERVER_2022), ("RHEL_9_STIG", "Red Hat Enterprise Linux 9 STIG")])
def test_a_checklist_stig_that_cannot_name_the_scan_does_not(tmp_path, stig_id, title):
    # No STIG ID (an empty ID would match every scan), or another STIG.
    rules = [{"group_id": "V-1", "rule_id_src": "SV-1r1_rule", "check_content": "check", "fix_text": "fix"}]
    reference = _checklist(tmp_path / "other.cklb", [{"stig_id": stig_id, "display_name": title, "version": 1,
                                                      "rules": rules}])
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [reference], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {""}
    assert result.coverage == {("WIN-SERVER-04", "")}


def test_a_reference_benchmark_with_no_id_names_no_scan_but_still_fills_by_rule(tmp_path):
    # An empty ID is a substring of every ID: it must not make one reference title every scan.
    nameless = tmp_path / "nameless.xml"
    nameless.write_text(
        '<Benchmark xmlns="http://checklists.nist.gov/xccdf/1.1" id=""><title>Some other STIG</title>'
        '<version>1</version><Group id="V-254239"><Rule id="SV-254239r945408_rule" severity="high">'
        "<version>WN22-00-000010</version><title>r</title><fixtext>fix</fixtext>"
        "<check><check-content>check</check-content></check></Rule></Group></Benchmark>", encoding="utf-8")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [nameless], tmp_path / "x")
    by_rule = {f.rule_id.rsplit("_rule_", 1)[-1]: f for f in result.findings}
    assert by_rule["SV-254241r945414_rule"].stig_title == ""            # not named by a reference with no ID
    assert ("WIN-SERVER-04", "") in result.coverage
    filled = by_rule["SV-254239r945408_rule"]                           # its own rule is still found, and labelled
    assert (filled.check_text, filled.stig_title) == ("check", "Some other STIG")
    assert filled.text_source == "Check and fix: nameless.xml V1"


def test_an_ignored_empty_checklist_names_nothing(tmp_path):
    empty = _checklist(tmp_path / "empty.cklb", [{"stig_id": "MS_Windows_Server_2022_STIG",
                                                  "display_name": SERVER_2022, "version": 1, "rules": []}])
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [empty], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {""}
    assert "empty.cklb: reference checklist contains 0 rules — ignored" in result.warnings


# --- which STIG of a checklist lacks check text -------------------------------------------

def _no_check_rules(prefix):
    return [{"group_id": f"V-{prefix}{n}", "rule_id_src": f"SV-{prefix}{n}r1_rule", "fix_text": "fix"} for n in (1, 2)]


def test_the_no_check_text_line_names_the_stig_when_the_checklist_holds_several(tmp_path):
    full = [{"group_id": "V-31", "rule_id_src": "SV-31r1_rule", "check_content": "check", "fix_text": "fix"}]
    reference = _checklist(tmp_path / "three.cklb", [
        {"stig_id": "A_STIG", "display_name": "A STIG", "version": 1, "rules": _no_check_rules(1)},
        {"stig_id": "B_STIG", "display_name": "B" * 500, "version": 1, "rules": _no_check_rules(2)},
        {"stig_id": "C_STIG", "display_name": "C STIG", "version": 1, "rules": full},
        {"stig_id": "D_STIG", "version": 1, "rules": _no_check_rules(4)[:1]},
    ])
    result = parse_stage([FIX / "scc_embedded_results.xml"], [reference], tmp_path / "x")
    lines = [w for w in result.warnings if "carry check text" in w]
    assert lines == [
        "three.cklb: A STIG: 0 of 2 rules carry check text",
        f"three.cklb: {'B' * 120}: 0 of 2 rules carry check text",
        "three.cklb: D_STIG: 0 of 1 rules carry check text",
    ]


def test_the_no_check_text_line_of_a_single_stig_checklist_names_the_file_only(tmp_path):
    reference = _checklist(tmp_path / "one.cklb", [
        {"stig_id": "A_STIG", "display_name": "A STIG", "version": 1, "rules": _no_check_rules(1)}])
    result = parse_stage([FIX / "scc_embedded_results.xml"], [reference], tmp_path / "x")
    assert "one.cklb: 0 of 2 rules carry check text" in result.warnings


# --- a ZIP is a folder: results inside one are never lost ---------------------------------

def _zip_of(path: Path, members: dict) -> Path:
    with zipfile.ZipFile(path, "w") as zf:
        for name, source in members.items():
            zf.write(source, name)
    return path


def test_a_checklist_zipped_into_the_results_list_is_a_scan(tmp_path):
    z = _zip_of(tmp_path / "host_checklists.zip", {"WIN-SERVER-04.cklb": FIX / "evaluate_stig_checklist.cklb"})
    result = parse_stage([FIX / "scc_embedded_results.xml", z], [], tmp_path / "x")
    assert result.source_file_count == 2
    from_checklist = [f for f in result.findings if f.server == "WIN-SERVER-01"]
    assert {f.vuln_id for f in from_checklist} == {"V-254239", "V-254241", "V-254242"}
    assert ("WIN-SERVER-01", "Microsoft Windows Server 2022 STIG") in result.coverage
    assert len([f for f in result.findings if f.server == "WKSTN-01"]) == 3
    assert not any("host_checklists.zip" in w for w in result.warnings)


@pytest.mark.parametrize("slot", ["results", "reference"])
def test_results_and_their_manual_stig_in_one_zip_are_both_read(tmp_path, slot):
    z = _zip_of(tmp_path / "scan_folder.zip", {
        "WKSTN-01_SCC-5.14_XCCDF-Results_MS_Windows_11.xml": FIX / "scc_embedded_results.xml",
        "U_MS_Windows_11_STIG_V2R9_Manual-xccdf.xml": FIX / "manual_stig_win11.xml",
    })
    result = parse_stage([z], [], tmp_path / "x") if slot == "results" else parse_stage([], [z], tmp_path / "x")
    assert result.source_file_count == 1 and len(result.findings) == 3
    assert ("WKSTN-01", "Microsoft Windows 11 STIG SCAP Benchmark") in result.coverage
    assert {f.vuln_id for f in result.findings if f.check_text} == {"V-253284", "V-253285"}   # the Manual was used too
    assert [s.file_name for s in result.enrichment.standalone_sources] == ["U_MS_Windows_11_STIG_V2R9_Manual-xccdf.xml"]
    assert not any("scan_folder.zip" in w for w in result.warnings)


# --- a file supplied twice is read once -----------------------------------------------------

def test_results_also_given_as_a_reference_are_one_scan(tmp_path):
    # `--references` pointed at the results folder does this.
    scc = FIX / "scc_embedded_results.xml"
    result = parse_stage([scc], [scc], tmp_path)
    assert (len(result.findings), result.source_file_count) == (3, 1)
    assert not any("identical" in w for w in result.warnings)
    nessus = FIX / "nessus_compliance.nessus"
    result = parse_stage([nessus], [nessus], tmp_path)
    assert (len(result.findings), result.source_file_count) == (4, 1)


def test_two_copies_of_one_scan_under_different_names_are_reported_once(tmp_path):
    first, second = tmp_path / "host.xml", tmp_path / "host - Copy.xml"
    for path in (first, second):
        path.write_bytes((FIX / "scc_embedded_results.xml").read_bytes())
    result = parse_stage([first, second], [], tmp_path / "x")
    assert (len(result.findings), result.source_file_count) == (3, 1)
    assert [w for w in result.warnings if "identical" in w] == ["host - Copy.xml: identical to host.xml — read once"]


def test_the_same_manual_stig_twice_is_one_reference_source(tmp_path):
    manual = FIX / "manual_stig_win11.xml"
    copy = tmp_path / "manual_again.xml"
    copy.write_bytes(manual.read_bytes())
    result = parse_stage([FIX / "scc_embedded_results.xml"], [manual, manual, copy], tmp_path / "x")
    assert [s.file_name for s in result.enrichment.standalone_sources] == ["manual_stig_win11.xml"]
    assert [w for w in result.warnings if "identical" in w] == [
        "manual_again.xml: identical to manual_stig_win11.xml — read once"]


# --- a scan is named only by a benchmark with the same ID -------------------------------------

NO_BENCHMARK = "no matching STIG benchmark"


def test_a_checklist_whose_stig_id_is_a_substring_does_not_name_the_scan(tmp_path):
    # "STIG" is a substring of every DISA benchmark ID. Only the row whose rule the checklist
    # really holds is titled, by enrichment; the other two stay untitled and are reported.
    reference = _one_rule_checklist(tmp_path / "other.cklb", stig_id="STIG", title="Some Other Product STIG")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [reference], tmp_path / "x")
    titles = {f.rule_id.rsplit("_rule_", 1)[-1]: f.stig_title for f in result.findings}
    assert titles == {"SV-254239r945408_rule": "Some Other Product STIG",
                      "SV-254241r945414_rule": "", "SV-254243r945420_rule": ""}
    assert ("WIN-SERVER-04", "") in result.coverage
    assert [w for w in result.warnings if NO_BENCHMARK in w] == [
        "evaluate_stig_results.xml: no matching STIG benchmark — 2 finding(s) have no STIG title, "
        "severity, or check/fix text. Add the STIG as a reference."]


def test_a_manual_stig_whose_id_is_a_substring_does_not_name_the_scan(tmp_path):
    other = tmp_path / "win11_shortid.xml"
    text = (FIX / "manual_stig_win11.xml").read_text(encoding="utf-8")
    assert 'id="MS_Windows_11_STIG"' in text
    other.write_text(text.replace('id="MS_Windows_11_STIG"', 'id="Windows"'), encoding="utf-8")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [other], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {""}
    assert result.coverage == {("WIN-SERVER-04", "")}
    assert len([w for w in result.warnings if NO_BENCHMARK in w and "3 finding(s)" in w]) == 1


def _evaluate_results_named(path: Path, benchmark_element: str) -> Path:
    text = (FIX / "evaluate_stig_results.xml").read_text(encoding="utf-8")
    start, end = text.index("<benchmark href="), text.index("/>", text.index("<benchmark href=")) + 2
    path.write_text(text[:start] + benchmark_element + text[end:], encoding="utf-8")
    return path


def test_a_scan_with_a_generic_href_and_no_id_is_not_named_by_any_reference(tmp_path):
    scan = _evaluate_results_named(tmp_path / "generic.xml", '<benchmark href="benchmark.xml"/>')
    for reference in (FIX / "scap_datastream_win11.xml", FIX / "sample_benchmark.xml"):
        result = parse_stage([scan], [reference], tmp_path / reference.stem)
        untitled = [f for f in result.findings if not f.check_text]
        assert all(f.stig_title == "" for f in untitled)
        if untitled:
            assert ("WIN-SERVER-04", "") in result.coverage
            assert [w for w in result.warnings if NO_BENCHMARK in w], result.warnings
        else:
            # Not named, yet every row was titled by its own rule: the scan covers those
            # titles, and no blank pair is left (see the coverage tests at the end).
            assert result.coverage == {(f.server, f.stig_title) for f in result.findings}
            assert all(f.stig_title for f in result.findings)


def test_a_longer_benchmark_id_is_a_different_benchmark(tmp_path):
    scan = _evaluate_results_named(
        tmp_path / "dc.xml",
        '<benchmark href="x.xml" id="xccdf_mil.disa.stig_benchmark_MS_Windows_Server_2022_STIG_DC"/>')
    result = parse_stage([scan], [FIX / "manual_stig_server2022.xml"], tmp_path / "x")
    assert ("WIN-SERVER-04", "") in result.coverage
    assert all(f.stig_title == "" for f in result.findings if not f.check_text)


def test_an_href_only_scan_is_read_against_the_one_benchmark_its_file_embeds(tmp_path):
    text = (FIX / "scc_embedded_results.xml").read_text(encoding="utf-8")
    element = ('<cdf:benchmark href="U_MS_Windows_11_V2R8_STIG_SCAP_1-3_Benchmark.xml" '
               'id="xccdf_mil.disa.stig_benchmark_Microsoft_Windows_11_STIG"/>')
    assert element in text
    scan = tmp_path / "href_only.xml"
    scan.write_text(text.replace(element, '<cdf:benchmark href="U_MS_Windows_11_V2R8_STIG_SCAP_1-3_Benchmark.xml"/>'),
                    encoding="utf-8")
    result = parse_stage([scan], [], tmp_path / "x")
    assert len(result.findings) == 3
    for f in result.findings:
        assert f.fix_text and f.vuln_id and f.severity.startswith("CAT ")
        assert (f.stig_title, f.scan_release) == ("Microsoft Windows 11 STIG SCAP Benchmark", "V2R8")
    assert not any(NO_BENCHMARK in w for w in result.warnings)


# --- a results checklist with values that are not text -----------------------------------

def _odd_results_checklist(path: Path) -> Path:
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    rules = doc["stigs"][0]["rules"]
    assert [r["status"] for r in rules[:3]] == ["open", "not_a_finding", "not_reviewed"]
    rules[0]["status"] = ["open"]
    rules[1]["status"] = 1
    rules[2]["check_content"] = {"x": 1}
    path.write_text(json.dumps(doc), encoding="utf-8")
    return path


def test_a_results_checklist_rule_with_an_unreadable_status_is_kept_for_review(tmp_path):
    plain = parse_stage([FIX / "evaluate_stig_checklist.cklb"], [], tmp_path / "a")
    odd = parse_stage([_odd_results_checklist(tmp_path / "odd_status.cklb")], [], tmp_path / "b")
    assert {f.vuln_id: f.status for f in plain.findings} == {
        "V-254239": "Open", "V-254241": "Not Reviewed", "V-254242": "Open"}
    assert {f.vuln_id: f.status for f in odd.findings} == {
        "V-254239": "Unknown", "V-254240": "Unknown", "V-254241": "Not Reviewed", "V-254242": "Open"}
    assert len(odd.findings) >= len(plain.findings)            # no row is lost


def test_non_text_values_in_a_results_checklist_are_counted_for_the_operator(tmp_path):
    odd = parse_stage([_odd_results_checklist(tmp_path / "odd_status.cklb")], [], tmp_path / "x")
    # The parser ignored 3 (two statuses and a check text), the reference pass 1 (it reads no
    # status): one line, and "at least", since the two passes read different values.
    assert [w for w in odd.warnings if "non-text" in w] == [
        "odd_status.cklb: at least 3 non-text value(s) in the checklist were ignored"]
    clean = parse_stage([FIX / "evaluate_stig_checklist.cklb"], [], tmp_path / "y")
    assert not any("non-text" in w for w in clean.warnings)
    not_reviewed = next(f for f in odd.findings if f.vuln_id == "V-254241")
    assert "{" not in not_reviewed.check_text


# --- a bad archive or an unreadable file is reported, never fatal ---------------------------

GOOD = FIX / "scc_embedded_results.xml"


def _good_scan_is_reported(result) -> bool:
    return len([f for f in result.findings if f.server == "WKSTN-01"]) == 3


def _set_encrypted_flag(zip_path: Path) -> None:
    data = bytearray(zip_path.read_bytes())
    for signature, flag_offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
        at = data.find(signature)
        while at != -1:
            data[at + flag_offset] |= 0x01
            at = data.find(signature, at + 4)
    zip_path.write_bytes(bytes(data))


@pytest.mark.parametrize("member", ["rhel9_results_2026-01-15T10:00:00.xml", "r" * 296 + ".xml"])
def test_a_member_name_that_is_not_a_valid_file_name_is_still_read(tmp_path, member):
    z = _zip_of(tmp_path / "scans.zip", {member: FIX / "openscap_results.xml"})
    result = parse_stage([GOOD, z], [], tmp_path / "x")
    assert _good_scan_is_reported(result) and result.source_file_count == 2
    assert any(f.server != "WKSTN-01" for f in result.findings)        # the member's own findings are there
    assert not any("scans.zip" in w or "Could not" in w for w in result.warnings)


def test_a_zipped_file_is_called_by_its_own_name_everywhere(tmp_path):
    # Text Source, the reference sources, per-file warnings and the host-name fallback all use
    # the name the member has in the archive, never the generated name it was extracted under.
    checklist = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    del checklist["target_data"]
    hostless = tmp_path / "hostless.cklb"
    hostless.write_text(json.dumps(checklist), encoding="utf-8")
    empty = tmp_path / "all_pass_none.xml"
    empty.write_text("<TestResult><target>H</target></TestResult>", encoding="utf-8")
    z = _zip_of(tmp_path / "bundle.zip", {
        "refs/U_MS_Windows_11_STIG_V2R9_Manual-xccdf.xml": FIX / "manual_stig_win11.xml",
        "lists/WIN-SERVER-09.cklb": hostless,
        "scans/no_results.xml": empty,
        "scans/evaluate.xml": FIX / "evaluate_stig_results.xml",
    })
    result = parse_stage([GOOD, z], [], tmp_path / "x")
    assert "Check: U_MS_Windows_11_STIG_V2R9_Manual-xccdf.xml V2R9 | Fix: scanner" in {f.text_source for f in result.findings}
    assert [s.file_name for s in result.enrichment.sources if not s.embedded] == ["U_MS_Windows_11_STIG_V2R9_Manual-xccdf.xml"]
    assert "WIN-SERVER-09" in {f.server for f in result.findings}       # the checklist names no host: its file name does
    assert any(w.startswith("no_results.xml: 0 rule results") for w in result.warnings), result.warnings
    assert any("WIN-SERVER-09.cklb V1R4" in f.text_source for f in result.findings if f.server == "WIN-SERVER-04")


def test_a_password_protected_archive_is_reported_and_the_run_continues(tmp_path):
    z = _zip_of(tmp_path / "locked.zip", {"scan.xml": FIX / "openscap_results.xml"})
    _set_encrypted_flag(z)
    result = parse_stage([GOOD, z], [], tmp_path / "x")
    assert _good_scan_is_reported(result) and result.source_file_count == 1
    assert [w for w in result.warnings if "locked.zip" in w] == ["locked.zip: password-protected — not read"]


def test_a_damaged_member_is_reported_and_the_run_continues(tmp_path):
    z = tmp_path / "damaged.zip"
    marker = b"MARKER-" + bytes(range(256))
    with zipfile.ZipFile(z, "w", zipfile.ZIP_STORED) as zf:
        zf.writestr("damaged.xml", b"<TestResult><!--" + marker + b"--></TestResult>")
        zf.write(FIX / "openscap_results.xml", "intact.xml")
    data = bytearray(z.read_bytes())
    data[data.find(marker) + 100] ^= 0xFF
    z.write_bytes(bytes(data))
    result = parse_stage([GOOD, z], [], tmp_path / "x")
    assert _good_scan_is_reported(result) and result.source_file_count == 2      # the intact member is a scan too
    assert [w for w in result.warnings if "damaged.zip" in w] == ["damaged.zip: could not read damaged.xml: Bad CRC-32 for file 'damaged.xml' — skipped"]


def test_an_archive_that_is_not_a_zip_is_reported_and_the_run_continues(tmp_path):
    bad = tmp_path / "truncated.zip"
    bad.write_bytes((FIX / "manual_stig_win11.xml").read_bytes()[:200])
    result = parse_stage([GOOD], [bad], tmp_path / "x")
    assert _good_scan_is_reported(result)
    assert [w for w in result.warnings if "truncated.zip" in w] == ["truncated.zip: not a readable ZIP — not read"]


@pytest.mark.parametrize("slot", ["results", "reference"])
def test_a_loose_file_that_cannot_be_read_is_reported_and_the_run_continues(tmp_path, slot):
    gone = tmp_path / "gone.xml"
    result = parse_stage([GOOD, gone], [], tmp_path / "x") if slot == "results" else parse_stage([GOOD], [gone], tmp_path / "x")
    assert _good_scan_is_reported(result) and result.source_file_count == 1
    assert [w for w in result.warnings if "gone.xml" in w] == ["Could not read file: gone.xml"]


@pytest.mark.parametrize("target,lost", [
    # The methods the pipeline calls (read / read_all: the result and why there is none).
    ("app.parsers.xccdf_parser.XCCDFResultsParser.read", ["openscap_results.xml"]),
    # read_all reads the reference; parse_all, the results file again for the benchmark it may embed
    ("app.parsers.benchmark_parser.BenchmarkParser.read_all", ["manual_stig_server2022.xml", "openscap_results.xml"]),
    ("app.parsers.nessus_parser.NessusComplianceParser.read", ["nessus_compliance.nessus"]),
])
def test_a_file_that_becomes_unreadable_while_it_is_parsed_does_not_stop_the_run(tmp_path, monkeypatch, target, lost):
    # The file was readable when it was classified; the read fails later (a share drops, a lock).
    module_name, class_name, method = target.rsplit(".", 2)
    import importlib
    cls = getattr(importlib.import_module(module_name), class_name)
    real = getattr(cls, method)

    def flaky(self, path, *args, **kwargs):
        if path.name in ("openscap_results.xml", "manual_stig_server2022.xml", "nessus_compliance.nessus"):
            raise PermissionError(13, "Permission denied", str(path))
        return real(self, path, *args, **kwargs)

    monkeypatch.setattr(cls, method, flaky)
    result = parse_stage([GOOD, FIX / "openscap_results.xml", FIX / "nessus_compliance.nessus"],
                         [FIX / "manual_stig_server2022.xml"], tmp_path / "x")
    assert _good_scan_is_reported(result)
    unreadable = [w for w in result.warnings if w.startswith("Could not read file: ")]
    assert unreadable == [f"Could not read file: {name}" for name in lost], result.warnings
    assert not any(f.server == "rhel9-lab-01" for f in result.findings) or "openscap_results.xml" not in lost


# --- what an archive skips reaches the operator ---------------------------------------------

def test_a_scan_nested_too_deep_and_a_broken_inner_zip_are_named(tmp_path):
    def wrap(inner: Path, name: str) -> Path:
        outer = tmp_path / name
        with zipfile.ZipFile(outer, "w") as zf:
            zf.write(inner, "inner.zip")
        return outer

    level = _zip_of(tmp_path / "l3.zip", {"deep_scan.xml": FIX / "openscap_results.xml"})
    for n in (2, 1):
        level = wrap(level, f"l{n}.zip")
    session = tmp_path / "session.zip"
    with zipfile.ZipFile(session, "w") as zf:
        zf.write(GOOD, "WKSTN-01_results.xml")
        zf.write(level, "archive_of_archives.zip")
        zf.writestr("corrupt.zip", b"not a zip at all")
    result = parse_stage([session], [], tmp_path / "x")
    assert _good_scan_is_reported(result) and result.source_file_count == 1
    assert [w for w in result.warnings if w.startswith("session.zip: ")] == [
        "session.zip: inner.zip is nested more than 2 ZIPs deep — skipped",
        "session.zip: corrupt.zip is not a readable ZIP — skipped",
    ]


def test_a_member_over_the_size_cap_reaches_the_operator(tmp_path, monkeypatch):
    import app.utils.zip_extract as zip_extract
    monkeypatch.setattr(zip_extract, "_MAX_EXTRACTED_BYTES", (FIX / "openscap_results.xml").stat().st_size - 1)
    session = _zip_of(tmp_path / "session.zip", {"big_scan.xml": FIX / "openscap_results.xml"})
    result = parse_stage([GOOD, session], [], tmp_path / "x")
    assert _good_scan_is_reported(result) and result.source_file_count == 1
    lines = [w for w in result.warnings if "session.zip" in w]
    assert len(lines) == 1 and lines[0].startswith("session.zip: big_scan.xml is larger than the ")
    assert lines[0].endswith("limit — skipped")


def test_more_than_five_skips_in_one_archive_are_counted(tmp_path):
    session = tmp_path / "session.zip"
    with zipfile.ZipFile(session, "w") as zf:
        zf.write(GOOD, "WKSTN-01_results.xml")
        for n in range(1, 9):
            zf.writestr(f"corrupt{n}.zip", b"not a zip at all %d" % n)
    result = parse_stage([session], [], tmp_path / "x")
    assert _good_scan_is_reported(result)
    assert [w for w in result.warnings if w.startswith("session.zip: ")] == [
        f"session.zip: corrupt{n}.zip is not a readable ZIP — skipped" for n in range(1, 6)] + [
        "session.zip: … and 3 more skipped"]


# --- the same scan in two encodings is counted once -------------------------------------------

def _arf_of(source: Path, target: Path) -> Path:
    """An ARF report collection wrapping the whole of *source*."""
    text = source.read_text(encoding="utf-8").split("?>", 1)[1]
    target.write_text(
        '<?xml version="1.0"?>\n<arf:asset-report-collection '
        'xmlns:arf="http://scap.nist.gov/schema/asset-reporting-format/1.1"><arf:reports><arf:report id="r1">'
        f"<arf:content>{text}</arf:content></arf:report></arf:reports></arf:asset-report-collection>",
        encoding="utf-8")
    return target


SAME_SCAN = "— its rows are not repeated"


def test_xccdf_results_and_the_arf_of_the_same_scan_are_one_scan(tmp_path):
    arf = _arf_of(GOOD, tmp_path / "arf.xml")
    session = _zip_of(tmp_path / "session.zip", {"WKSTN-01_XCCDF-Results.xml": GOOD, "WKSTN-01_ARF-Results.xml": arf})
    result = parse_stage([session], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"WKSTN-01_ARF-Results.xml: same host and results as WKSTN-01_XCCDF-Results.xml {SAME_SCAN}"]
    assert all(f.text_source == "Check: no reference supplied | Fix: scanner" for f in result.findings)
    assert result.coverage == {("WKSTN-01", "Microsoft Windows 11 STIG SCAP Benchmark")}


def test_the_first_file_wins_whichever_encoding_it_is(tmp_path):
    arf = _arf_of(GOOD, tmp_path / "WKSTN-01_ARF.xml")
    result = parse_stage([arf, GOOD], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"scc_embedded_results.xml: same host and results as WKSTN-01_ARF.xml {SAME_SCAN}"]


def test_two_scans_of_one_host_and_benchmark_with_different_results_are_both_kept(tmp_path):
    text = GOOD.read_text(encoding="utf-8")
    rescan = tmp_path / "rescan.xml"
    changed = text.replace('idref="xccdf_mil.disa.stig_rule_SV-253284r958928_rule" severity="high"><cdf:result>fail',
                           'idref="xccdf_mil.disa.stig_rule_SV-253284r958928_rule" severity="high"><cdf:result>pass')
    assert changed != text
    rescan.write_text(changed, encoding="utf-8")
    result = parse_stage([GOOD, rescan], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 5)
    assert not any("same host and results as" in w for w in result.warnings)


def test_scans_with_no_rule_results_are_never_merged(tmp_path):
    empties = []
    for name in ("empty_a.xml", "empty_b.xml"):
        path = tmp_path / name
        path.write_text(f"<TestResult><!-- {name} --><target>HOST-E</target>"
                        "<target-address>10.0.0.9</target-address></TestResult>", encoding="utf-8")
        empties.append(path)
    result = parse_stage([GOOD, *empties], [], tmp_path / "x")
    assert not any("same host and results as" in w for w in result.warnings)
    assert [w for w in result.warnings if "0 rule results" in w] == [
        f"{name}: 0 rule results — not counted as a scan; if this file is a benchmark, pass it as a reference"
        for name in ("empty_a.xml", "empty_b.xml")]


@pytest.mark.parametrize("bare_first", [True, False])
def test_of_two_copies_of_a_scan_the_one_carrying_its_benchmark_is_kept(tmp_path, bare_first):
    # A bare TestResult and the full SCC file of the same scan: the full copy is read, whichever comes
    # first, and the bare copy leaves no trace (no untitled coverage pair, no text labelled as a reference).
    text = GOOD.read_text(encoding="utf-8")
    start, end = text.index("<cdf:TestResult"), text.index("</cdf:TestResult>") + len("</cdf:TestResult>")
    bare = tmp_path / "bare_results.xml"
    bare.write_text('<?xml version="1.0"?>\n' + text[start:end].replace(
        "<cdf:TestResult", '<cdf:TestResult xmlns:cdf="http://checklists.nist.gov/xccdf/1.2"', 1), encoding="utf-8")
    result = parse_stage([bare, GOOD] if bare_first else [GOOD, bare], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"bare_results.xml: same host and results as scc_embedded_results.xml {SAME_SCAN}"]
    for f in result.findings:
        assert f.fix_text and f.severity.startswith("CAT ") and f.scan_release == "V2R8"
        assert f.text_source == "Check: no reference supplied | Fix: scanner"
    assert result.coverage == {("WKSTN-01", "Microsoft Windows 11 STIG SCAP Benchmark")}
    assert [s.file_name for s in result.enrichment.sources] == ["scc_embedded_results.xml"]


EVAL = FIX / "evaluate_stig_results.xml"
SERVER_2022 = "Microsoft Windows Server 2022 STIG"


def _checklist_of_eval_host(path: Path, *, indent: int | None = None, **target) -> Path:
    """The fixture checklist moved to WIN-SERVER-04, the host of evaluate_stig_results.xml."""
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    doc["target_data"].update({"host_name": "WIN-SERVER-04", "ip_address": "192.168.1.40", **target})
    path.write_text(json.dumps(doc, indent=indent), encoding="utf-8")
    return path


def _xccdf_of_the_checklists_run(path: Path) -> Path:
    """evaluate_stig_results.xml with the results the checklist holds, as when Evaluate-STIG writes
    both files from one run (the two fixtures were made separately and disagree on two rules)."""
    text = EVAL.read_text(encoding="utf-8")
    changed = text.replace("<result>unknown</result>", "<result>notapplicable</result>").replace(
        "<score ", '<rule-result idref="xccdf_mil.disa.stig_rule_SV-254242r958490_rule" severity="low">'
                   "<result>fail</result></rule-result>\n  <score ")
    assert changed.count("<result>") == text.count("<result>") + 1
    path.write_text(changed, encoding="utf-8")
    return path


@pytest.mark.parametrize("xccdf_first", [True, False])
def test_the_xccdf_results_and_the_checklist_of_one_run_are_one_scan(tmp_path, xccdf_first):
    xccdf = _xccdf_of_the_checklists_run(tmp_path / "WIN-SERVER-04_XCCDF.xml")
    checklist = _checklist_of_eval_host(tmp_path / "WIN-SERVER-04.cklb")
    result = parse_stage([xccdf, checklist] if xccdf_first else [checklist, xccdf], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    # The checklist is kept: its rows carry check and fix text, the XCCDF rows none.
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"WIN-SERVER-04_XCCDF.xml: same host and results as WIN-SERVER-04.cklb {SAME_SCAN}"]
    assert all(f.text_source == "Check and fix: scanner" for f in result.findings)
    assert result.coverage == {("WIN-SERVER-04", SERVER_2022)}
    assert not any(TWICE in w for w in result.warnings)


def test_a_checklist_saved_twice_with_different_formatting_is_one_scan(tmp_path):
    first = _checklist_of_eval_host(tmp_path / "WIN-SERVER-04.cklb")
    resaved = _checklist_of_eval_host(tmp_path / "WIN-SERVER-04_resaved.cklb", indent=2)
    assert first.read_bytes() != resaved.read_bytes()
    result = parse_stage([first, resaved], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"WIN-SERVER-04_resaved.cklb: same host and results as WIN-SERVER-04.cklb {SAME_SCAN}"]
    assert [s.file_name for s in result.enrichment.sources] == ["WIN-SERVER-04.cklb"]


# --- rules reported by more than one file are named ------------------------------------------------

TWICE = "— each file's rows are in the report, so totals count those rules more than once"


def _win11_checklist(path: Path, statuses: dict[int, str], host: str = "WKSTN-01") -> Path:
    """A Windows 11 STIG checklist for *host*: one rule per entry of *statuses* (rule number -> status)."""
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    doc["target_data"].update(host_name=host, ip_address="10.0.0.21")
    stig = doc["stigs"][0]
    stig.update(display_name="Microsoft Windows 11 STIG", stig_id="MS_Windows_11_STIG",
                stig_name="Microsoft Windows 11 Security Technical Implementation Guide")
    template = stig["rules"][0]
    stig["rules"] = [dict(template, rule_id=f"SV-{n}r1_rule", rule_id_src=f"SV-{n}r1_rule", group_id=f"V-{n}",
                          group_id_src=f"V-{n}", rule_version=_W11_STIG_IDS.get(n, f"WN11-{n}"), status=status)
                     for n, status in statuses.items()]
    path.write_text(json.dumps(doc), encoding="utf-8")
    return path


def test_a_scan_and_a_checklist_of_one_host_that_disagree_name_the_rules_counted_twice(tmp_path):
    # The SCC scan fails 253284, 253285 and 253286; the checklist, titled differently, agrees on two.
    checklist = _win11_checklist(tmp_path / "WKSTN-01_Win11.cklb",
                                 {253284: "open", 253285: "open", 253286: "not_a_finding", 253290: "open"})
    result = parse_stage([GOOD, checklist], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 6)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 2 rule(s) appear in more than one of 2 files (scc_embedded_results.xml, WKSTN-01_Win11.cklb) {TWICE}"]


def test_a_scan_titled_by_a_reference_and_its_checklist_name_the_rules_counted_twice(tmp_path):
    # The XCCDF rows take the Manual STIG's title, the checklist rows the checklist's: not one title.
    checklist = _checklist_of_eval_host(tmp_path / "WIN-SERVER-04.cklb")
    result = parse_stage([EVAL, checklist], [FIX / "manual_stig_server2022.xml"], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 6)
    assert len({f.stig_title for f in result.findings}) == 2
    assert not any("same host and results as" in w for w in result.warnings)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WIN-SERVER-04: 2 rule(s) appear in more than one of 2 files (evaluate_stig_results.xml, WIN-SERVER-04.cklb) {TWICE}"]


def test_one_host_and_stig_reported_by_two_files_with_different_results_are_both_kept_and_named(tmp_path):
    # The two fixtures as they are: one host and STIG, but not the same results.
    checklist = _checklist_of_eval_host(tmp_path / "WIN-SERVER-04.cklb")
    result = parse_stage([EVAL, checklist], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 6)
    assert not any("same host and results as" in w for w in result.warnings)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WIN-SERVER-04: 2 rule(s) appear in more than one of 2 files (evaluate_stig_results.xml, WIN-SERVER-04.cklb) {TWICE}"]


def test_two_files_covering_different_rules_of_one_host_say_nothing(tmp_path):
    checklist = _win11_checklist(tmp_path / "WKSTN-01_extra.cklb", {253284: "not_a_finding", 253290: "open"})
    result = parse_stage([GOOD, checklist], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 4)
    assert not any(TWICE in w for w in result.warnings)


def test_two_hosts_with_the_same_failing_rules_are_not_merged(tmp_path):
    other = tmp_path / "WKSTN-02.xml"
    other.write_text(GOOD.read_text(encoding="utf-8").replace("WKSTN-01", "WKSTN-02"), encoding="utf-8")
    result = parse_stage([GOOD, other], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 6)
    assert {f.server for f in result.findings} == {"WKSTN-01", "WKSTN-02"}
    assert not any("same host and results as" in w or TWICE in w for w in result.warnings)


def test_at_most_five_hosts_are_named_and_the_rest_counted(tmp_path):
    # Eight hosts, each scanned twice with different results: 253290 fails in January only.
    text = GOOD.read_text(encoding="utf-8")
    january = text.replace('SV-253290r958940_rule" severity="medium"><cdf:result>pass',
                           'SV-253290r958940_rule" severity="medium"><cdf:result>fail')
    assert january != text
    paths = []
    for i in range(8):
        for month, body in (("jan", january), ("mar", text)):
            path = tmp_path / f"WKSTN-{i:02d}_{month}.xml"
            path.write_text(body.replace("WKSTN-01", f"WKSTN-{i:02d}"), encoding="utf-8")
            paths.append(path)
    result = parse_stage(paths, [], tmp_path / "x")
    assert result.source_file_count == 16
    lines = [w for w in result.warnings if TWICE in w or "more host(s)" in w]
    assert lines == [
        f"WKSTN-{i:02d}: 3 rule(s) appear in more than one of 2 files (WKSTN-{i:02d}_jan.xml, WKSTN-{i:02d}_mar.xml) {TWICE}"
        for i in range(5)] + ["… and 3 more host(s) with rules reported by more than one file"]


def test_the_host_and_files_named_are_escaped_and_the_files_bounded(tmp_path):
    # Six checklists of one host, each from another address: six scans, one line naming five files.
    paths = [_checklist_of_eval_host(tmp_path / f"host_{i}.cklb", host_name="EVIL\nHOST", ip_address=f"10.0.0.{i}")
             for i in range(6)]
    result = parse_stage(paths, [], tmp_path / "x")
    assert result.source_file_count == 6
    line, = [w for w in result.warnings if TWICE in w]
    assert "\n" not in line
    assert line == ("EVIL\\nHOST: 3 rule(s) appear in more than one of 6 files "
                    f"(host_0.cklb, host_1.cklb, host_2.cklb, host_3.cklb, host_4.cklb …) {TWICE}")


# --- rules a results checklist gives no status for ---------------------------------------------

def test_rules_skipped_for_a_blank_status_reach_the_operator(tmp_path):
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    rules = doc["stigs"][0]["rules"]
    rules[0]["status"] = ""             # was open
    rules[1]["status"] = None           # was not_a_finding
    del rules[2]["status"]              # was not_reviewed
    blank = tmp_path / "blank_status.cklb"
    blank.write_text(json.dumps(doc), encoding="utf-8")
    result = parse_stage([blank], [], tmp_path / "x")
    assert [f.vuln_id for f in result.findings] == ["V-254242"]
    assert "blank_status.cklb: 3 rule(s) have no status and were skipped" in result.warnings
    clean = parse_stage([FIX / "evaluate_stig_checklist.cklb"], [], tmp_path / "y")
    assert not any("no status" in w for w in clean.warnings)


def test_a_zipped_checklist_with_skipped_rules_is_named_by_its_own_name(tmp_path):
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    doc["stigs"][0]["rules"][1]["status"] = ""
    blank = tmp_path / "blank.cklb"
    blank.write_text(json.dumps(doc), encoding="utf-8")
    result = parse_stage([_zip_of(tmp_path / "lists.zip", {"WIN-SERVER-01.cklb": blank})], [], tmp_path / "x")
    assert "WIN-SERVER-01.cklb: 1 rule(s) have no status and were skipped" in result.warnings


# --- an ordinary long file name is shown whole --------------------------------------------------

LONG_SCC_NAME = "WIN-SERVER-04_SCC-5.14.1_2026-01-15_100000_XCCDF-Results_MS_Windows_Server_2022_STIG-002.003.xml"


@pytest.mark.parametrize("zipped", [False, True])
def test_a_long_scc_file_name_is_not_cut_in_warnings(tmp_path, zipped):
    assert len(LONG_SCC_NAME) > 80
    scan = tmp_path / LONG_SCC_NAME
    scan.write_bytes((FIX / "evaluate_stig_results.xml").read_bytes())
    supplied = _zip_of(tmp_path / "session.zip", {LONG_SCC_NAME: scan}) if zipped else scan
    result = parse_stage([supplied], [], tmp_path / "x")
    assert (f"{LONG_SCC_NAME}: no matching STIG benchmark — 3 finding(s) have no STIG title, severity, "
            "or check/fix text. Add the STIG as a reference.") in result.warnings


# --- a parse error names the file, never a path on this server ----------------------------------

@pytest.mark.parametrize("as_reference", [False, True])
def test_a_parse_error_names_the_file_and_no_path_on_this_server(tmp_path, caplog, as_reference):
    broken = "<TestResult><a></TestResult>"       # a damaged scan: classified by its tags, named by its parser
    loose = tmp_path / "broken_loose.xml"
    loose.write_text(broken, encoding="utf-8")
    archive = tmp_path / "session.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("scans/broken_member.xml", broken.replace("<a>", "<b>"))   # not the loose file's bytes
        zf.writestr("scans/broken.nessus", "<NessusClientData_v2><a></NessusClientData_v2>")
        zf.writestr("scans/broken.cklb", "{not json")
        zf.writestr("scans/inner.zip", b"not a zip")
    extract = tmp_path / "x"
    supplied = [loose, archive]
    with caplog.at_level(logging.DEBUG, logger="app"):
        result = parse_stage([GOOD] + ([] if as_reference else supplied), supplied if as_reference else [], extract)
    lines = [record.getMessage() for record in caplog.records] + result.warnings
    server = ["member_", "file:/"]
    for directory in (extract, tmp_path, Path(tempfile.gettempdir())):
        server += [str(directory).lower(), directory.as_posix().lower()]
    leaks = [line for line in lines if any(form in line.lower() for form in server)]
    assert leaks == []
    # The operator still learns which file is broken, and how.
    for name in ("broken_loose.xml", "broken_member.xml", "broken.nessus"):
        assert any(name in line and "Opening and ending tag mismatch" in line for line in lines), name
    assert any("inner.zip" in line and "not a readable ZIP" in line for line in lines)


# --- a warning names the file it is about ------------------------------------------------------

def test_each_warning_names_the_file_it_is_about_when_names_are_shared(tmp_path):
    # A loose results.xml with no rule results, and a ZIP holding a good scan and another empty
    # one, both called results.xml: each warning must name the file it is about, and only names
    # the operator supplied.
    def empty(path: Path, host: str) -> Path:
        path.write_text(f"<TestResult><target>{host}</target><target-address>10.0.0.9</target-address></TestResult>",
                        encoding="utf-8")
        return path

    loose = empty(tmp_path / "results.xml", "HOST-L")
    z = _zip_of(tmp_path / "more.zip", {"x/results.xml": GOOD, "y/results.xml": empty(tmp_path / "y.xml", "HOST-Y")})
    result = parse_stage([loose, z], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (3, 3)
    assert [w for w in result.warnings if "0 rule results" in w] == [
        f"{name}: 0 rule results — not counted as a scan; if this file is a benchmark, pass it as a reference"
        for name in ("results.xml", "more.zip/y/results.xml")]
    assert not any("results_1" in w for w in result.warnings)



# --- malformed checklist entries reach the operator -----------------------------------------------

def _checklist_with(tmp_path: Path, name: str, extra_rules: list, extra_stigs: list) -> Path:
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    doc["stigs"][0]["rules"] += extra_rules
    doc["stigs"] += extra_stigs
    path = tmp_path / name
    path.write_text(json.dumps(doc), encoding="utf-8")
    return path


def test_malformed_entries_of_a_results_checklist_reach_the_operator(tmp_path):
    many = _checklist_with(tmp_path, "many.cklb", [["x"], {"status": "open"}], [42, {"display_name": "No rules"}])
    one = _checklist_with(tmp_path, "one.cklb", [], ["not a stig"])
    doc = json.loads(one.read_text(encoding="utf-8"))
    doc["target_data"]["host_name"] = "OTHER-HOST"          # not the same scan as many.cklb
    one.write_text(json.dumps(doc), encoding="utf-8")
    result = parse_stage([many, one], [], tmp_path / "x")
    assert [w for w in result.warnings if "malformed" in w] == [
        "many.cklb: 4 malformed checklist entries were skipped",
        "one.cklb: 1 malformed checklist entry was skipped"]
    clean = parse_stage([FIX / "evaluate_stig_checklist.cklb"], [], tmp_path / "y")
    assert not any("malformed" in w for w in clean.warnings)


def test_malformed_entries_of_a_reference_checklist_reach_the_operator(tmp_path):
    reference = _checklist_with(tmp_path, "reference.cklb", ["x", 7], [{"display_name": "No rules", "rules": 3}])
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [reference], tmp_path / "x")
    assert [w for w in result.warnings if "malformed" in w] == [
        "reference.cklb: 3 malformed checklist entries were skipped"]


def test_a_reference_checklist_with_only_malformed_entries_says_both(tmp_path):
    path = tmp_path / "empty_ref.cklb"
    path.write_text(json.dumps({"stigs": ["x", {"stig_id": "A"}]}), encoding="utf-8")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [path], tmp_path / "x")
    assert "empty_ref.cklb: 2 malformed checklist entries were skipped" in result.warnings
    assert "empty_ref.cklb: reference checklist contains 0 rules — ignored" in result.warnings



# --- only copies of one scan are merged --------------------------------------------------------------

_W11_STIG_IDS = {253284: "WN11-00-000150", 253285: "WN11-00-000160", 253286: "WN11-00-000170", 253290: "WN11-00-000200"}


def _nessus(path: Path, hosts: dict[tuple[str, str], list[tuple[int, str, str]]], *, pretty: bool = False,
            rule_ids: bool = True, vuln_ids: bool = True, stig_ids: dict[int, str] | None = None,
            benchmark: str = "DISA Microsoft Windows 11 STIG", revision: int = 1) -> Path:
    """A .nessus compliance scan: (host name, IP) -> [(rule number, CAT, result)], Windows 11 STIG.

    Without *rule_ids* each item carries its STIG ID (from *stig_ids*, else the Windows 11 one) and,
    unless *vuln_ids* is false, its V-ID, as older DISA audit files do."""
    stig_ids = stig_ids or _W11_STIG_IDS
    sep = "\n  " if pretty else ""
    body = ""
    for (name, ip), items in hosts.items():
        body += (f'<ReportHost name="{ip}">{sep}<HostProperties><tag name="host-ip">{ip}</tag>'
                 f'<tag name="hostname">{name}</tag></HostProperties>')
        for number, cat, outcome in items:
            body += (f'{sep}<ReportItem pluginName="Windows Compliance Checks" pluginFamily="Policy Compliance">'
                     f"<cm:compliance-check-name>check {number}</cm:compliance-check-name>"
                     f"<cm:compliance-result>{outcome}</cm:compliance-result>"
                     f"<cm:compliance-reference>CAT|{cat},"
                     + (f"Rule-ID|SV-{number}r{revision}_rule," if rule_ids else f"STIG-ID|{stig_ids[number]},")
                     + (f"Vuln-ID|V-{number}" if vuln_ids else "") + "</cm:compliance-reference>"
                     f"<cm:compliance-benchmark-name>{benchmark}</cm:compliance-benchmark-name>"
                     "<cm:compliance-solution>Fix it.</cm:compliance-solution></ReportItem>")
        body += "</ReportHost>"
    path.write_text('<?xml version="1.0"?><NessusClientData_v2><Report name="r" xmlns:cm="http://www.nessus.org/cm">'
                    f"{body}</Report></NessusClientData_v2>", encoding="utf-8")
    return path


def test_a_scan_and_a_nessus_scan_reporting_the_same_rows_are_not_merged(tmp_path):
    # The .nessus file's WKSTN-01 rows match the SCC scan row for row; its other host passed every
    # check, so it adds no actionable row. They are still two scans by two scanners.
    weekly = _nessus(tmp_path / "weekly.nessus", {
        ("WKSTN-01", "10.0.0.21"): [(253284, "I", "FAILED"), (253285, "II", "FAILED"), (253286, "III", "FAILED")],
        ("WKSTN-99", "10.0.0.99"): [(253284, "I", "PASSED")]})
    result = parse_stage([GOOD, weekly], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 6)
    assert not any("same host and results as" in w for w in result.warnings)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 3 rule(s) appear in more than one of 2 files (scc_embedded_results.xml, weekly.nessus) {TWICE}"]
    scc_rows = [f for f in result.findings if f.stig_title == "Microsoft Windows 11 STIG SCAP Benchmark"]
    assert len(scc_rows) == 3
    for f in scc_rows:
        assert f.scan_release == "V2R8" and f.text_source == "Check: no reference supplied | Fix: scanner"


def test_two_copies_of_one_nessus_scan_are_merged(tmp_path):
    hosts = {("WKSTN-01", "10.0.0.21"): [(253284, "I", "FAILED"), (253285, "II", "FAILED")]}
    first = _nessus(tmp_path / "weekly.nessus", hosts)
    again = _nessus(tmp_path / "weekly_export.nessus", hosts, pretty=True)
    assert first.read_bytes() != again.read_bytes()
    result = parse_stage([first, again], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 2)
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"weekly_export.nessus: same host and results as weekly.nessus {SAME_SCAN}"]


def test_two_files_with_one_result_set_but_titles_of_different_products_are_not_merged(tmp_path):
    statuses = {253284: "open", 253285: "open"}
    windows = _win11_checklist(tmp_path / "windows.cklb", statuses)
    other = _win11_checklist(tmp_path / "other.cklb", statuses)
    doc = json.loads(other.read_text(encoding="utf-8"))
    doc["stigs"][0]["display_name"] = "Red Hat Enterprise Linux 9 STIG"
    other.write_text(json.dumps(doc), encoding="utf-8")
    result = parse_stage([windows, other], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 4)
    assert not any("same host and results as" in w for w in result.warnings)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 2 rule(s) appear in more than one of 2 files (windows.cklb, other.cklb) {TWICE}"]


def test_copies_whose_titles_name_one_product_in_other_words_are_still_merged(tmp_path):
    # With the Manual STIG as reference the XCCDF rows are titled "… Security Technical Implementation
    # Guide" and the checklist rows "… STIG": one product, one scan.
    xccdf = _xccdf_of_the_checklists_run(tmp_path / "WIN-SERVER-04_XCCDF.xml")
    checklist = _checklist_of_eval_host(tmp_path / "WIN-SERVER-04.cklb")
    result = parse_stage([xccdf, checklist], [FIX / "manual_stig_server2022.xml"], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"WIN-SERVER-04_XCCDF.xml: same host and results as WIN-SERVER-04.cklb {SAME_SCAN}"]



# --- a copy that is not read still lends its STIG text -------------------------------------------------

def test_a_copy_that_is_not_read_fills_the_blanks_the_kept_copy_left(tmp_path):
    # Copy A carries check and fix text for two rules and nothing for the third (4); copy B, the
    # same scan, fix text for all three (3). A is read; B's fix text still reaches the third rule.
    text = GOOD.read_text(encoding="utf-8")
    copy_a = text
    for number in ("253284", "253285"):
        ref = f'name="oval:mil.disa.stig.windows11:def:{number}"/>'
        copy_a = copy_a.replace(ref, ref + f"<cdf:check-content>Verify {number}.</cdf:check-content>")
    copy_a = copy_a.replace(
        '<cdf:fixtext fixref="F-56739r958931_fix">Configure the policy value for Configure SMB v1 client driver '
        "to Disabled.</cdf:fixtext>", "")
    assert copy_a.count("<cdf:check-content>") == 2 and copy_a.count("<cdf:fixtext") == text.count("<cdf:fixtext") - 1
    a = tmp_path / "copy_A.xml"
    a.write_text(copy_a, encoding="utf-8")
    b = tmp_path / "copy_B_arf.xml"
    b.write_text(text, encoding="utf-8")
    result = parse_stage([a, b], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert [w for w in result.warnings if "same host and results as" in w] == [f"copy_B_arf.xml: same host and results as copy_A.xml {SAME_SCAN}"]
    by_rule = {f.vuln_id: f for f in result.findings}
    assert by_rule["V-253284"].text_source == by_rule["V-253285"].text_source == "Check and fix: scanner"
    third = by_rule["V-253286"]
    assert third.fix_text.startswith("Configure the policy value")
    assert third.text_source == "Check: no reference supplied | Fix: copy_B_arf.xml V2R8"
    assert not any("neither check nor fix text" in w for w in result.warnings)


@pytest.mark.parametrize("xccdf_first", [True, False])
def test_an_identical_copy_that_is_not_read_is_never_cited(tmp_path, xccdf_first):
    arf = _arf_of(GOOD, tmp_path / "WKSTN-01_ARF.xml")
    result = parse_stage([GOOD, arf] if xccdf_first else [arf, GOOD], [], tmp_path / "x")
    kept = "scc_embedded_results.xml" if xccdf_first else "WKSTN-01_ARF.xml"
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert all(f.text_source == "Check: no reference supplied | Fix: scanner" for f in result.findings)
    assert [s.file_name for s in result.enrichment.sources] == [kept]



# --- a long display name keeps its file name ---------------------------------------------------------

def test_a_long_display_name_in_text_source_keeps_its_file_name(tmp_path):
    folder = "STIG_Library_April_2026/Operating_Systems/Microsoft/Windows/Workstation_and_Server_Benchmarks_Manual_Editions"
    library = _zip_of(tmp_path / "lib.zip", {f"{folder}/win11/manual-xccdf.xml": FIX / "manual_stig_win11.xml",
                                             f"{folder}/srv22/manual-xccdf.xml": FIX / "manual_stig_server2022.xml"})
    result = parse_stage([GOOD], [library], tmp_path / "x")
    shown = f"lib.zip/{folder}/win11/manual-xccdf.xml"
    assert len(shown) > 120
    filled = [f for f in result.findings if f.check_text]
    assert filled and all(f.text_source.startswith("Check: …" + shown[-119:] + " V2R9") for f in filled)



# --- the overlap line matches a rule by any of its IDs ------------------------------------------------

def test_rows_naming_one_rule_by_different_ids_are_found_in_both_files(tmp_path):
    # The .nessus items carry the STIG ID and V-ID but no rule ID: keyed by its STIG ID, the Nessus
    # row never met the SCC row keyed by its DISA stem. Both name the rule by V-ID and STIG ID.
    audit = _nessus(tmp_path / "old_audit.nessus",
                    {("WKSTN-01", "10.0.0.21"): [(253284, "I", "FAILED"), (253285, "II", "FAILED")]}, rule_ids=False)
    result = parse_stage([GOOD, audit], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 5)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 2 rule(s) appear in more than one of 2 files (scc_embedded_results.xml, old_audit.nessus) {TWICE}"]


def test_rules_shared_by_different_pairs_of_files_count_the_files_involved(tmp_path):
    # a and b share 253284, b and c share 253285, a and c share 253286: three rules, three files.
    a = _win11_checklist(tmp_path / "a.cklb", {253284: "open", 253286: "open"})
    b = _win11_checklist(tmp_path / "b.cklb", {253284: "open", 253285: "open"})
    c = _win11_checklist(tmp_path / "c.cklb", {253285: "open", 253286: "open"})
    result = parse_stage([a, b, c], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (3, 6)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 3 rule(s) appear in more than one of 3 files (a.cklb, b.cklb, c.cklb) {TWICE}"]


def test_two_files_of_one_host_that_share_no_id_say_nothing(tmp_path):
    audit = _nessus(tmp_path / "other_rule.nessus", {("WKSTN-01", "10.0.0.21"): [(253290, "II", "FAILED")]},
                    rule_ids=False)
    result = parse_stage([GOOD, audit], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 4)
    assert not any(TWICE in w for w in result.warnings)



# --- a file with no XML element near its start says so -------------------------------------------

_LICENSE = '<?xml version="1.0"?>\n<!--' + "license text " * 6000 + "-->\n"      # 78 KB before the root


def test_a_real_scan_behind_a_long_comment_is_named_for_what_it_lacks_not_called_empty(tmp_path):
    padded = tmp_path / "WKSTN-77_results.xml"
    padded.write_text(_LICENSE + GOOD.read_text(encoding="utf-8").split("?>", 1)[1].replace("WKSTN-01", "WKSTN-77"),
                      encoding="utf-8")
    result = parse_stage([GOOD, padded], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert "WKSTN-77_results.xml: no XML element in its first 64 KiB — not read" in result.warnings
    assert not any("0 rule results" in w for w in result.warnings)
    zipped = _zip_of(tmp_path / "padded.zip", {"WKSTN-78_results.xml": padded})
    result = parse_stage([GOOD, zipped], [], tmp_path / "y")
    assert (result.source_file_count, len(result.findings)) == (1, 3)
    assert ("padded.zip: 1 file(s) have no XML element in their first 64 KiB — not read: WKSTN-78_results.xml"
            in result.warnings)


def test_a_manual_stig_behind_a_long_comment_is_named_for_what_it_lacks(tmp_path):
    manual = tmp_path / "manual_padded.xml"
    manual.write_text('<?xml version="1.0"?>\n<!--' + "x" * 70000 + "-->\n"
                      + (FIX / "manual_stig_server2022.xml").read_text(encoding="utf-8").split("?>", 1)[1], encoding="utf-8")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [manual], tmp_path / "x")
    assert "manual_padded.xml: no XML element in its first 64 KiB — not read" in result.warnings
    assert not any("Could not parse benchmark" in w for w in result.warnings)



# --- a shared STIG ID does not make two identified rules one ----------------------------------------

def _ndm_checklist(path: Path, product: str, number: int) -> Path:
    """A one-rule Cisco NDM checklist of CORE-SW-01: router and switch NDM share STIG ID CISC-ND-000010."""
    doc = json.loads((FIX / "evaluate_stig_checklist.cklb").read_text(encoding="utf-8-sig"))
    doc["target_data"].update(host_name="CORE-SW-01", ip_address="10.9.9.9")
    stig = doc["stigs"][0]
    stig.update(display_name=f"Cisco IOS XE {product} NDM STIG", stig_name=f"Cisco IOS XE {product} NDM STIG",
                stig_id=f"Cisco_IOS-XE_{product}_NDM_STIG")
    stig["rules"] = [dict(stig["rules"][0], rule_id=f"SV-{number}r960735_rule", rule_id_src=f"SV-{number}r960735_rule",
                          group_id=f"V-{number}", group_id_src=f"V-{number}", rule_version="CISC-ND-000010",
                          status="open")]
    path.write_text(json.dumps(doc), encoding="utf-8")
    return path


def test_two_rules_sharing_only_a_stig_id_are_two_rules(tmp_path):
    router = _ndm_checklist(tmp_path / "CORE-SW-01_Router_NDM.cklb", "Router", 215807)
    switch = _ndm_checklist(tmp_path / "CORE-SW-01_Switch_NDM.cklb", "Switch", 220518)
    result = parse_stage([router, switch], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 2)
    assert not any(TWICE in w for w in result.warnings)


def test_a_row_known_only_by_its_stig_id_joins_the_one_rule_with_that_stig_id(tmp_path):
    audit = _nessus(tmp_path / "stig_only.nessus", {("WKSTN-01", "10.0.0.21"): [(253284, "I", "FAILED")]},
                    rule_ids=False, vuln_ids=False)
    result = parse_stage([GOOD, audit], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 4)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 1 rule(s) appear in more than one of 2 files (scc_embedded_results.xml, stig_only.nessus) {TWICE}"]


def test_a_row_known_only_by_a_stig_id_two_rules_share_is_not_attributed(tmp_path):
    # Its title names neither NDM STIG, so nothing can tell which of the two rules it reports.
    router = _ndm_checklist(tmp_path / "CORE-SW-01_Router_NDM.cklb", "Router", 215807)
    switch = _ndm_checklist(tmp_path / "CORE-SW-01_Switch_NDM.cklb", "Switch", 220518)
    audit = _nessus(tmp_path / "stig_only.nessus", {("CORE-SW-01", "10.9.9.9"): [(215807, "II", "FAILED")]},
                    rule_ids=False, vuln_ids=False, stig_ids={215807: "CISC-ND-000010"})
    result = parse_stage([router, switch, audit], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (3, 3)
    assert not any(TWICE in w for w in result.warnings)


def test_rules_with_their_own_stems_are_not_one_rule_because_a_checklist_repeats_a_stig_id(tmp_path):
    checklist = _win11_checklist(tmp_path / "WKSTN-01_Win11.cklb", {253284: "open", 253285: "open"})
    doc = json.loads(checklist.read_text(encoding="utf-8"))
    for rule in doc["stigs"][0]["rules"]:
        rule["rule_version"] = "WN11-X"
    checklist.write_text(json.dumps(doc), encoding="utf-8")
    result = parse_stage([GOOD, checklist], [], tmp_path / "x")
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 2 rule(s) appear in more than one of 2 files (scc_embedded_results.xml, WKSTN-01_Win11.cklb) {TWICE}"]



def test_parse_stage_cancellation_reaches_inside_an_archive(tmp_path):
    z = _zip_of(tmp_path / "many.zip", {f"WKSTN-{i}.xml": GOOD for i in range(3)})
    calls = []

    class Cancelled(Exception):
        pass

    def check() -> None:
        calls.append(1)
        if len(calls) == 3:            # parse_stage's own first poll, the archive, then its first member
            raise Cancelled

    with pytest.raises(Cancelled):
        parse_stage([z], [], tmp_path / "x", cancel_check=check)
    assert not list((tmp_path / "x").iterdir())       # nothing was extracted before the third poll



# --- pins for behaviour the mutation run found unpinned ------------------------------------------------

def test_a_row_known_only_by_a_stig_id_joins_its_rule_even_when_enrichment_cannot_name_it(tmp_path):
    # The audit's title names no product the SCC benchmark matches, so enrichment fills no V-ID:
    # the row stays known only by its STIG ID and must join the SCC row through it.
    audit = _nessus(tmp_path / "site_audit.nessus", {("WKSTN-01", "10.0.0.21"): [(253284, "I", "FAILED")]},
                    rule_ids=False, vuln_ids=False, benchmark="Site Audit")
    result = parse_stage([GOOD, audit], [], tmp_path / "x")
    audit_row, = [f for f in result.findings if f.stig_title == "Site Audit"]
    assert audit_row.vuln_id == ""
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-01: 1 rule(s) appear in more than one of 2 files (scc_embedded_results.xml, site_audit.nessus) {TWICE}"]


def test_rows_known_only_by_one_stig_id_join_each_other_when_no_identified_rule_has_it(tmp_path):
    a = _nessus(tmp_path / "a.nessus", {("WKSTN-50", "10.0.0.50"): [(253284, "I", "FAILED"), (253285, "II", "FAILED")]},
                rule_ids=False, vuln_ids=False)
    b = _nessus(tmp_path / "b.nessus", {("WKSTN-50", "10.0.0.50"): [(253284, "I", "FAILED")]},
                rule_ids=False, vuln_ids=False)
    result = parse_stage([a, b], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 3)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-50: 1 rule(s) appear in more than one of 2 files (a.nessus, b.nessus) {TWICE}"]


def test_one_rule_at_two_revisions_joins_through_its_stem_alone(tmp_path):
    # No V-ID and no STIG ID: only the DISA stem (SV-253284) says r1 and r2 are one rule.
    a = _nessus(tmp_path / "a.nessus", {("WKSTN-50", "10.0.0.50"): [(253284, "I", "FAILED")]}, vuln_ids=False)
    b = _nessus(tmp_path / "b.nessus", {("WKSTN-50", "10.0.0.50"): [(253284, "I", "FAILED"), (253285, "II", "FAILED")]},
                vuln_ids=False, revision=2)
    result = parse_stage([a, b], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 3)
    assert {f.rule_id for f in result.findings} == {"SV-253284r1_rule", "SV-253284r2_rule", "SV-253285r2_rule"}
    assert not any(f.vuln_id or f.stig_id for f in result.findings)
    assert [w for w in result.warnings if TWICE in w] == [
        f"WKSTN-50: 1 rule(s) appear in more than one of 2 files (a.nessus, b.nessus) {TWICE}"]


def test_a_stig_id_standing_in_for_a_rule_id_does_not_join_two_identified_rules(tmp_path):
    # Items with a V-ID and no Rule-ID: Nessus puts the STIG-ID where the rule ID goes.
    router = _nessus(tmp_path / "router.nessus", {("CORE-SW-01", "10.9.9.9"): [(215807, "II", "FAILED")]},
                     rule_ids=False, stig_ids={215807: "CISC-ND-000010"}, benchmark="Cisco IOS XE Router NDM STIG")
    switch = _nessus(tmp_path / "switch.nessus", {("CORE-SW-01", "10.9.9.9"): [(220518, "II", "FAILED")]},
                     rule_ids=False, stig_ids={220518: "CISC-ND-000010"}, benchmark="Cisco IOS XE Switch NDM STIG")
    result = parse_stage([router, switch], [], tmp_path / "x")
    assert (result.source_file_count, len(result.findings)) == (2, 2)
    assert {f.rule_id for f in result.findings} == {"CISC-ND-000010"}
    assert not any(TWICE in w for w in result.warnings)


def test_every_cancellation_poll_of_parse_stage_is_reached(tmp_path):
    checklist = _win11_checklist(tmp_path / "WKSTN-77.cklb", {253284: "open"}, host="WKSTN-77")
    supplied = ([GOOD, checklist], [FIX / "manual_stig_win11.xml", FIX / "evaluate_stig_checklist.cklb"])

    class Cancelled(Exception):
        pass

    def run(stop_at: int | None) -> int:
        calls = []

        def check() -> None:
            calls.append(1)
            if len(calls) == stop_at:
                raise Cancelled
        if stop_at is None:
            parse_stage(*supplied, tmp_path / f"x{stop_at}", cancel_check=check)
        else:
            with pytest.raises(Cancelled):
                parse_stage(*supplied, tmp_path / f"x{stop_at}", cancel_check=check)
        return len(calls)

    # One poll to start, one per supplied file while classifying, one per reference benchmark, reference
    # checklist, XCCDF results file and results checklist while parsing, and one before matching.
    assert run(None) == 1 + 4 + 1 + 1 + 1 + 1 + 1
    for stop_at in range(1, 11):
        assert run(stop_at) == stop_at


def test_a_title_filled_by_enrichment_is_in_the_coverage(tmp_path):
    # The reference's benchmark ID is not the scan's, so it does not name the scan: only enrichment
    # titles its rows, and the coverage must carry that title.
    reference = tmp_path / "manual_other_id.xml"
    reference.write_text((FIX / "manual_stig_server2022.xml").read_text(encoding="utf-8").replace(
        'id="MS_Windows_Server_2022_STIG"', 'id="Some_Other_Benchmark"', 1), encoding="utf-8")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [reference], tmp_path / "x")
    title = "Microsoft Windows Server 2022 Security Technical Implementation Guide"
    assert any(f.stig_title == title for f in result.findings)
    assert ("WIN-SERVER-04", title) in result.coverage


def test_an_untitled_reference_never_names_a_scan(tmp_path):
    def benchmark(path: Path, title: str) -> Path:
        path.write_text(
            '<Benchmark xmlns="http://checklists.nist.gov/xccdf/1.1" id="MS_Windows_Server_2022_STIG">'
            f"{title}<Group id='V-999999'><Rule id='SV-999999r1_rule' severity='low'><version>X-1</version>"
            "<fixtext>Fix.</fixtext></Rule></Group></Benchmark>", encoding="utf-8")
        return path
    untitled = benchmark(tmp_path / "untitled.xml", "")
    titled = benchmark(tmp_path / "titled.xml", "<title>Server 2022 Reference</title>")
    result = parse_stage([FIX / "evaluate_stig_results.xml"], [untitled, titled], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {"Server 2022 Reference"}


def test_on_a_tie_the_copy_supplied_first_is_kept_whatever_its_format(tmp_path):
    # The checklist carries fix text on the same three rows as the SCC file and no check text: a tie.
    checklist = _win11_checklist(tmp_path / "WKSTN-01_Win11.cklb", {253284: "open", 253285: "open", 253286: "open"})
    doc = json.loads(checklist.read_text(encoding="utf-8"))
    for rule in doc["stigs"][0]["rules"]:
        rule["check_content"] = ""
    checklist.write_text(json.dumps(doc), encoding="utf-8")
    result = parse_stage([checklist, GOOD], [], tmp_path / "x")
    assert [w for w in result.warnings if "same host and results as" in w] == [
        f"scc_embedded_results.xml: same host and results as WKSTN-01_Win11.cklb {SAME_SCAN}"]


def test_an_embedded_benchmark_with_no_rules_is_not_the_scans_own(tmp_path):
    # A results file whose Benchmark element holds no rules: the reference with the scan's ID names it.
    test_result = (FIX / "evaluate_stig_results.xml").read_text(encoding="utf-8").split("-->", 1)[1]
    results = tmp_path / "empty_benchmark_results.xml"
    results.write_text('<Benchmark xmlns="http://checklists.nist.gov/xccdf/1.2" '
                       'id="xccdf_mil.disa.stig_benchmark_MS_Windows_Server_2022_STIG"><title>Empty</title>'
                       + test_result.replace('xmlns="http://checklists.nist.gov/xccdf/1.2"', "", 1) + "</Benchmark>",
                       encoding="utf-8")
    result = parse_stage([results], [FIX / "manual_stig_server2022.xml"], tmp_path / "x")
    assert {f.stig_title for f in result.findings} == {
        "Microsoft Windows Server 2022 Security Technical Implementation Guide"}
    assert not any("not found in benchmark" in w for w in result.warnings)


def test_loose_scap_support_files_are_not_broken_benchmarks_or_source_files(tmp_path):
    # A loose OVAL or CPE file is what a ZIP's member of the same kind is: support content,
    # named once for the run. It is not "a benchmark that could not be parsed", not a results
    # file that failed, and not counted as a source file.
    oval = tmp_path / "U_MS_Windows_11_V2R8_STIG_SCAP_1-3_Benchmark-oval.xml"
    oval.write_text("<oval_definitions xmlns='http://oval.mitre.org/XMLSchema/oval-definitions-5'/>",
                    encoding="utf-8")
    cpe = tmp_path / "cpe.xml"
    cpe.write_text("<cpe-list xmlns='http://cpe.mitre.org/dictionary/2.0'/>", encoding="utf-8")
    result = parse_stage([FIX / "scc_embedded_results.xml", cpe], [FIX / "manual_stig_win11.xml", oval],
                         tmp_path / "x")
    assert not [w for w in result.warnings if "Could not parse" in w or "cpe.xml:" in w], result.warnings
    assert [w for w in result.warnings if "SCAP support" in w] == [
        "2 SCAP support file(s) (OVAL, CPE, OCIL, stylesheets) were not used: "
        "cpe.xml, U_MS_Windows_11_V2R8_STIG_SCAP_1-3_Benchmark-oval.xml"]
    assert result.source_file_count == 1
    assert result.findings[0].text_source == "Check: manual_stig_win11.xml V2R9 | Fix: scanner"


# --- a benchmark-less scan titled by enrichment: its coverage pair ---------------------------

WIN11_MANUAL = "Microsoft Windows 11 Security Technical Implementation Guide"


def _bare_scan(path: Path, passing: tuple[str, ...] = ()) -> Path:
    """WKSTN-01's SCC scan as a bare TestResult (no benchmark of its own; its ID is not the
    Manual STIG's), with the rules whose stems are in *passing* set to pass."""
    text = GOOD.read_text(encoding="utf-8")
    start, end = text.index("<cdf:TestResult"), text.index("</cdf:TestResult>") + len("</cdf:TestResult>")
    lines = text[start:end].replace(
        "<cdf:TestResult", '<cdf:TestResult xmlns:cdf="http://checklists.nist.gov/xccdf/1.2"', 1).splitlines()
    for stem in passing:
        (i,) = [n for n, line in enumerate(lines) if f"_rule_{stem}r" in line]
        lines[i] = lines[i].replace("<cdf:result>fail<", "<cdf:result>pass<")
    path.write_text('<?xml version="1.0"?>\n' + "\n".join(lines), encoding="utf-8")
    return path


def test_a_benchmark_less_scan_whose_rows_were_all_titled_covers_those_titles(tmp_path):
    # The one rule the Manual STIG lacks passes: every actionable row is titled by enrichment,
    # so the scan covers that title and no (host, "") pair is left to fail closed on.
    scan = _bare_scan(tmp_path / "bare.xml", passing=("SV-253286",))
    result = parse_stage([scan], [FIX / "manual_stig_win11.xml"], tmp_path / "x")
    assert [f.stig_title for f in result.findings] == [WIN11_MANUAL, WIN11_MANUAL]
    assert result.coverage == {("WKSTN-01", WIN11_MANUAL)}


def test_a_benchmark_less_scan_with_an_untitled_row_keeps_its_blank_pair(tmp_path):
    scan = _bare_scan(tmp_path / "bare.xml")            # SV-253286 fails, and no reference holds it
    result = parse_stage([scan], [FIX / "manual_stig_win11.xml"], tmp_path / "x")
    assert sorted(f.stig_title for f in result.findings) == ["", WIN11_MANUAL, WIN11_MANUAL]
    assert result.coverage == {("WKSTN-01", ""), ("WKSTN-01", WIN11_MANUAL)}


def test_an_all_pass_benchmark_less_scan_keeps_its_blank_pair(tmp_path):
    scan = _bare_scan(tmp_path / "bare.xml", passing=("SV-253284", "SV-253285", "SV-253286"))
    result = parse_stage([scan], [FIX / "manual_stig_win11.xml"], tmp_path / "x", allow_empty=True)
    assert result.findings == [] and result.coverage == {("WKSTN-01", "")}


def test_the_delta_of_scans_titled_by_their_reference_names_no_untitled_pair(tmp_path):
    from openpyxl import load_workbook

    from app.exporters.excel_exporter import ExcelExporter
    from app.processors.delta import compute_delta
    refs = [FIX / "manual_stig_win11.xml"]
    base = parse_stage([_bare_scan(tmp_path / "base.xml", passing=("SV-253286",))], refs, tmp_path / "b")
    curr = parse_stage([_bare_scan(tmp_path / "curr.xml", passing=("SV-253286", "SV-253285"))], refs, tmp_path / "c")
    delta = compute_delta(base.findings, curr.findings,
                          baseline_coverage=base.coverage, current_coverage=curr.coverage)
    assert {f.vuln_id: f.delta_status for f in delta.findings} == {"V-253284": "Persisting", "V-253285": "Resolved"}
    assert (delta.warnings, delta.not_rescanned_pairs, delta.newly_scanned_pairs) == ([], set(), set())
    out = tmp_path / "d.xlsx"
    ExcelExporter().export_delta(delta, out)
    cells = [c.value for row in load_workbook(out)["Summary"].iter_rows() for c in row if c.value is not None]
    assert "(no STIG title)" not in cells and "Warnings" not in cells


# --- delta warnings name references and claim only what happened -----------------------------

def _delta_of(base, curr):
    from app.processors.delta import compute_delta
    return compute_delta(base.findings, curr.findings, baseline_coverage=base.coverage, current_coverage=curr.coverage)


def test_the_v_id_warning_states_what_was_measured_and_no_hint_it_cannot_support(tmp_path):
    # Both runs have the same reference. The baseline fails SV-254243, which the reference lacks
    # (so the finding has no V-ID); the current run passes it. Nothing says the runs were given
    # different references, so the warning must not suggest it.
    text = EVAL.read_text(encoding="utf-8")
    at = text.index("<result>unknown</result>", text.index("SV-254243r945420_rule"))
    current = tmp_path / "current.xml"
    current.write_text(text[:at] + "<result>pass</result>" + text[at + len("<result>unknown</result>"):],
                       encoding="utf-8")
    refs = [FIX / "manual_stig_server2022.xml"]
    delta = _delta_of(parse_stage([EVAL], refs, tmp_path / "b"), parse_stage([current], refs, tmp_path / "c"))
    assert sorted(f.delta_status for f in delta.findings) == ["Persisting", "Persisting", "Resolved"]
    assert [w for w in delta.warnings if "V-ID" in w] == [
        "On hosts in both runs, 1 of 3 baseline and 0 of 2 current finding(s) have no V-ID, "
        "so they were matched by rule ID instead."]
    assert not any("--benchmarks" in w or "only one run" in w for w in delta.warnings)


def test_a_pair_warning_claims_no_tagged_findings_when_none_were(tmp_path):
    # A bare baseline (its SV-253286 row untitled: blank pair kept) against the full SCC file of
    # the same scan: every row matches, Persisting. The blank pair is still listed, but no
    # finding was tagged Not re-scanned for it.
    refs = [FIX / "manual_stig_win11.xml"]
    delta = _delta_of(parse_stage([_bare_scan(tmp_path / "base.xml")], refs, tmp_path / "b"),
                      parse_stage([GOOD], refs, tmp_path / "c"))
    assert {f.delta_status for f in delta.findings} == {"Persisting"}
    assert delta.not_rescanned_pairs == {("WKSTN-01", "")}
    assert not [w for w in delta.warnings if "not re-scanned in the current set" in w], delta.warnings
    assert [w for w in delta.warnings if "cannot be verified as re-scanned, because" in w] == [
        "1 host/STIG pair(s) cannot be verified as re-scanned, because a run "
        "scanned the host with no STIG title: WKSTN-01 / (no STIG title)"]
    assert [w for w in delta.warnings if "no STIG title (" in w] == [
        "1 host(s) have scans with no STIG title (the scan could not be matched to a benchmark); "
        "no finding was tagged for it — supply the STIG as a reference for both runs: WKSTN-01"]
    assert not any("are tagged" in w or "--benchmarks" in w for w in delta.warnings), delta.warnings


def test_a_titled_pair_is_unverifiable_when_the_other_run_scanned_the_host_untitled(tmp_path):
    # The reviewer's D3. Baseline: bare scan, its rows titled by the reference except SV-253286.
    # Current: the same host, all pass, bare: nothing to title, so it covers only (host, "").
    # The host WAS re-scanned, so the titled pair is not "not re-scanned"; it cannot be verified,
    # like the blank pair. Statuses are unchanged (fail closed). The advice is what can help:
    # the current file needs its benchmark for its coverage to be named.
    refs = [FIX / "manual_stig_win11.xml"]
    passing = ("SV-253284", "SV-253285", "SV-253286")
    delta = _delta_of(parse_stage([_bare_scan(tmp_path / "base.xml")], refs, tmp_path / "b"),
                      parse_stage([_bare_scan(tmp_path / "curr.xml", passing)], refs, tmp_path / "c",
                                  allow_empty=True))
    assert [f.delta_status for f in delta.findings] == ["Not re-scanned"] * 3
    assert not [w for w in delta.warnings if "were not re-scanned" in w or "nobody re-ran" in w], delta.warnings
    assert [w for w in delta.warnings if "cannot be verified as re-scanned, because" in w] == [
        "2 host/STIG pair(s) cannot be verified as re-scanned, because a run "
        f"scanned the host with no STIG title: WKSTN-01 / (no STIG title); WKSTN-01 / {WIN11_MANUAL}"]
    assert [w for w in delta.warnings if "no STIG title (" in w] == [
        "1 host(s) have scans with no STIG title (the scan could not be matched to a benchmark), so 3 "
        "finding(s) there with no match in the other run cannot be verified as re-scanned and are tagged "
        "Not re-scanned / Newly scanned; a results file with no benchmark of its own is named only "
        "through rows a reference titles; to name its coverage, give it its benchmark (an SCC results "
        "file that embeds it, or the matching benchmark as a reference): WKSTN-01"]
    expected = {"not re-scanned": set(), "newly scanned": set(),
                "unverifiable": {"WKSTN-01 / (no STIG title)", f"WKSTN-01 / {WIN11_MANUAL}"}}
    assert _gaps_in_the_workbook(delta, tmp_path / "d3.xlsx") == expected
    assert _gaps_in_the_warnings(delta.warnings) == expected


def test_a_blank_pair_in_both_runs_is_called_unverifiable_not_unscanned(tmp_path):
    # The reviewer's D2: a bare scan in both runs, SV-253286 failing (untitled) in both and
    # SV-253290 failing in the current run only. The blank pair is in both coverages: the scan
    # was re-run, it only cannot be verified, so neither "not re-scanned" nor "no baseline scan".
    refs = [FIX / "manual_stig_win11.xml"]
    current = _bare_scan(tmp_path / "curr.xml")
    current.write_text(current.read_text(encoding="utf-8").replace(
        'SV-253290r958940_rule" severity="medium"><cdf:result>pass<',
        'SV-253290r958940_rule" severity="medium"><cdf:result>fail<'), encoding="utf-8")
    delta = _delta_of(parse_stage([_bare_scan(tmp_path / "base.xml")], refs, tmp_path / "b"),
                      parse_stage([current], refs, tmp_path / "c"))
    assert sorted(f.delta_status for f in delta.findings) == ["Newly scanned", "Persisting", "Persisting", "Persisting"]
    assert not [w for w in delta.warnings if "were not re-scanned" in w or "have no baseline scan" in w], delta.warnings
    assert [w for w in delta.warnings if "cannot be verified as re-scanned, because" in w] == [
        "1 host/STIG pair(s) cannot be verified as re-scanned, because a run "
        "scanned the host with no STIG title: WKSTN-01 / (no STIG title)"]


# --- the delta workbook's Coverage block says what the warning lines say ----------------------

_GAP_HEADINGS = {
    "not re-scanned": "Host / STIG pairs not re-scanned",
    "newly scanned": "Host / STIG pairs newly scanned",
    "unverifiable": "Host / STIG pairs that cannot be verified as re-scanned (a run has no STIG title for the host)",
}
_GAP_WARNINGS = {
    "not re-scanned": "baseline host/STIG pair(s) were not re-scanned in the current set",
    "newly scanned": "current host/STIG pair(s) have no baseline scan",
    "unverifiable": "host/STIG pair(s) cannot be verified as re-scanned, because a run scanned the host",
}


def _gaps_in_the_workbook(delta, path: Path) -> dict[str, set[str]]:
    """{meaning: "HOST / STIG" pairs listed under its Coverage heading} of the exported delta."""
    from openpyxl import load_workbook

    from app.exporters.excel_exporter import ExcelExporter
    ExcelExporter().export_delta(delta, path)
    rows = [[c.value for c in r] for r in load_workbook(path)["Summary"].iter_rows()]
    out = {}
    for meaning, heading in _GAP_HEADINGS.items():
        at = next(i for i, r in enumerate(rows) if r[0] == heading)
        out[meaning] = {f"{rows[at + 1 + k][1]} / {rows[at + 1 + k][2]}" for k in range(rows[at][1])}
    return out


def _gaps_in_the_warnings(warnings: list[str]) -> dict[str, set[str]]:
    """{meaning: "HOST / STIG" pairs its warning line names} (empty when there is no such line)."""
    out = {meaning: set() for meaning in _GAP_WARNINGS}
    for w in warnings:
        for meaning, phrase in _GAP_WARNINGS.items():
            if phrase in w:
                out[meaning] = set(w.rsplit(": ", 1)[1].split("; "))
    return out


def test_the_coverage_block_and_the_warnings_agree_on_the_reviewers_d2(tmp_path):
    # A bare scan in both runs: the blank pair was re-scanned but cannot be verified.
    refs = [FIX / "manual_stig_win11.xml"]
    current = _bare_scan(tmp_path / "curr.xml")
    current.write_text(current.read_text(encoding="utf-8").replace(
        'SV-253290r958940_rule" severity="medium"><cdf:result>pass<',
        'SV-253290r958940_rule" severity="medium"><cdf:result>fail<'), encoding="utf-8")
    delta = _delta_of(parse_stage([_bare_scan(tmp_path / "base.xml")], refs, tmp_path / "b"),
                      parse_stage([current], refs, tmp_path / "c"))
    expected = {"not re-scanned": set(), "newly scanned": set(),
                "unverifiable": {"WKSTN-01 / (no STIG title)"}}
    assert _gaps_in_the_workbook(delta, tmp_path / "d2.xlsx") == expected
    assert _gaps_in_the_warnings(delta.warnings) == expected


def test_the_coverage_block_and_the_warnings_agree_on_a_missing_host(tmp_path):
    # HOST-B is not in the current run at all (its untitled pair was not re-scanned). HOST-C was
    # re-scanned for another STIG only, all titled: its Edge STIG was not re-scanned, and its new
    # STIG is newly scanned. HOST-A was re-scanned untitled only: neither its untitled pair nor
    # its Edge STIG can be verified.
    from app.parsers.base import Finding
    from app.processors.delta import compute_delta

    def finding(vuln: str, host: str, title: str) -> Finding:
        return Finding(title, vuln, f"SV-{vuln[2:]}r1_rule", "CAT II", "Open", host, "10.0.0.1", "c", "f")

    base = [finding("V-1", "HOST-A", ""), finding("V-2", "HOST-B", ""), finding("V-3", "HOST-A", "Edge STIG"),
            finding("V-4", "HOST-C", "Edge STIG")]
    curr = [finding("V-1", "HOST-A", ""), finding("V-9", "HOST-C", "Win STIG")]
    delta = compute_delta(base, curr, baseline_coverage={(f.server, f.stig_title) for f in base},
                          current_coverage={(f.server, f.stig_title) for f in curr})
    expected = {"not re-scanned": {"HOST-B / (no STIG title)", "HOST-C / Edge STIG"},
                "newly scanned": {"HOST-C / Win STIG"},
                "unverifiable": {"HOST-A / (no STIG title)", "HOST-A / Edge STIG"}}
    assert _gaps_in_the_workbook(delta, tmp_path / "missing.xlsx") == expected
    assert _gaps_in_the_warnings(delta.warnings) == expected


# --- a file that was not read: ParseResult.warnings says why ----------------------------------------------

def test_the_run_says_why_each_file_was_not_read(tmp_path):
    from tests.test_web import broken_uploads, expected_reasons
    results, references = broken_uploads()
    (tmp_path / "in").mkdir()

    def write(name, payload):
        path = tmp_path / "in" / name
        path.write_bytes(payload)
        return path
    result = parse_stage([FIX / "scc_embedded_results.xml"] + [write(n, b) for n, b in results],
                         [write(n, b) for n, b in references], tmp_path / "x")
    for name, line in expected_reasons(tmp_path).items():
        assert [w for w in result.warnings if name in w] == [line], (name, result.warnings)


# --- every problem is reported once, by the run, with its reason (tests/problem_cases.py) --------------

from difflib import SequenceMatcher  # noqa: E402

from app.reference.normalize import escape_controls  # noqa: E402
from tests.problem_cases import GOOD as CASE_GOOD, problem_cases  # noqa: E402

_CASES = problem_cases()


def _run_case(case, tmp_path, monkeypatch):
    import app.utils.zip_extract as zip_extract
    for name, value in case.patches.items():
        monkeypatch.setattr(zip_extract, name, value)
    folder = tmp_path / "in"
    folder.mkdir()
    results, references = [CASE_GOOD], []
    for zone, name, payload in case.files:
        path = folder / name
        path.write_bytes(payload)
        (results if zone == "results" else references).append(path)
    return parse_stage(results, references, tmp_path / "x")


@pytest.mark.parametrize("case", _CASES, ids=[case.id for case in _CASES])
def test_the_run_reports_each_problem_once_with_its_reason(case, tmp_path, monkeypatch):
    result = _run_case(case, tmp_path, monkeypatch)
    assert [w for w in result.warnings if case.subject in w] == case.expected


# The modules that read uploads. What they log at WARNING while the run parses is detail the run
# does not report: a host or IP address taken from the file name, the TestResult used, a row
# left out (the run counts them), duplicate rule IDs. Every other problem they meet travels in
# what they return, and the run words it once. A new WARNING line in these modules must be one of
# these, or be returned instead.
_READERS = ("app.parsers", "app.reference.cklb_loader", "app.utils.zip_extract")
_DETAIL_TEMPLATES = {
    "%s: %d <TestResult> elements found (remediation-style output) — using the last one (post-remediation state)",
    "%s: No hostname found in <target>, <target-facts>, or <title> — using filename '%s'",
    "%s: No IP address found in <target-address> or <target-facts> — using 'N/A'",
    "%s: no host_name/fqdn in target_data — using filename '%s'",
    "%s: no ip_address in target_data — using 'N/A'",
    "Benchmark %s: %d duplicate rule ID(s) — the last definition of each is used: %s%s",
    # Per-row lines (RowLog: the first five of a file): which row; the run gives the count.
    "%s: rule-result '%s' has no <result> element — skipping",
    "%s: rule with no group_id/rule_id — skipping",
    "%s: rule %s has a status that is not text — recording as 'Unknown' so it is not silently dropped",
    "%s: rule %s has no status — skipping",
    "%s: rule %s has unrecognised status '%s' — recording as 'Unknown' so it is not silently dropped",
    "%s: STIG '%s' has no rules list — skipping",
    "%s: ReportHost with no name/fqdn — using filename '%s'",
    "%s: compliance item with unrecognised result '%s' — recording as 'Unknown' so it is not silently dropped",
}


def _normalised(text: str) -> str:
    return " ".join(escape_controls(text).split()).casefold()


@pytest.mark.parametrize("case", _CASES, ids=[case.id for case in _CASES])
def test_no_reader_log_line_repeats_a_reason_the_run_reports(case, tmp_path, monkeypatch, caplog):
    # Which lines a reader may log is checked statically below; this checks what they say, as
    # written for each problem, against what the run reports, so the web UI never shows it twice.
    with caplog.at_level(logging.WARNING, logger="app"):
        result = _run_case(case, tmp_path, monkeypatch)
    warnings = [_normalised(w) for w in result.warnings]
    for record in caplog.records:
        if record.levelno < logging.WARNING or not record.name.startswith(_READERS):
            continue
        message = _normalised(record.getMessage())
        for warning in warnings:
            shared = SequenceMatcher(None, message, warning, autojunk=False).find_longest_match(
                0, len(message), 0, len(warning))
            assert shared.size < 24, (record.getMessage(), warning)


def test_every_allowed_detail_line_is_one_a_reader_writes():
    # The list above names real lines: a reworded detail line must be listed again, not slip past.
    import ast
    import app
    root = Path(app.__file__).parent
    written = set()
    for module in ("parsers/xccdf_parser.py", "parsers/cklb_parser.py", "parsers/nessus_parser.py",
                   "parsers/benchmark_parser.py", "reference/cklb_loader.py", "utils/zip_extract.py"):
        tree = ast.parse((root / module).read_text(encoding="utf-8"))
        written |= {node.value for node in ast.walk(tree) if isinstance(node, ast.Constant) and isinstance(node.value, str)}
    assert sorted(_DETAIL_TEMPLATES - written) == []


# RowLog logs the row lines it is given: its own call forwards them, and each call of a row_log
# is checked instead.
_FORWARDERS = {"RowLog.__call__"}
_LOG_METHODS = {"warning", "error", "critical", "exception", "log"}


def _reader_log_calls(source: str) -> list[tuple[str, ast.Call]]:
    """(qualified enclosing function, call) for every logging call in *source*."""
    calls: list[tuple[str, ast.Call]] = []

    def visit(node, scope):
        for child in ast.iter_child_nodes(node):
            if isinstance(child, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef)):
                visit(child, [*scope, child.name])
                continue
            if isinstance(child, ast.Call):
                func = child.func
                if (isinstance(func, ast.Attribute) and func.attr in _LOG_METHODS) or (
                        isinstance(func, ast.Name) and func.id == "row_log"):
                    calls.append((".".join(scope), child))
            visit(child, scope)
    visit(ast.parse(source), [])
    return calls


def _stray_reader_log_calls(sources: dict[str, str]) -> list[str]:
    stray = []
    for module, source in sources.items():
        for scope, call in _reader_log_calls(source):
            if scope in _FORWARDERS:
                continue
            first = call.args[0] if call.args else None
            if not (isinstance(first, ast.Constant) and first.value in _DETAIL_TEMPLATES):
                stray.append(f"{module}:{call.lineno} in {scope or '<module>'}")
    return stray


def _reader_sources() -> dict[str, str]:
    import app
    root = Path(app.__file__).parent
    paths = [*sorted((root / "parsers").glob("*.py")), root / "reference" / "cklb_loader.py",
             root / "utils" / "zip_extract.py"]
    return {path.relative_to(root).as_posix(): path.read_text(encoding="utf-8") for path in paths}


def test_every_reader_log_call_is_a_known_detail_line():
    # The guard above sees only the branches a problem case reaches. This one reads every log
    # call in the reader modules: each must log a listed detail line.
    assert _stray_reader_log_calls(_reader_sources()) == []


def test_a_stray_log_call_in_an_unreached_branch_is_found():
    # The check itself: a WARNING no problem case reaches is still found.
    sources = _reader_sources()
    sources["parsers/cklb_parser.py"] += "\n\ndef _unreached():\n    log.warning('x')\n"
    assert [line.split(" in ")[1] for line in _stray_reader_log_calls(sources)] == ["_unreached"]


def test_a_reference_past_the_rule_limit_is_named_once_and_not_called_empty(tmp_path, monkeypatch):
    import app.reference.library as library_module
    monkeypatch.setattr(library_module, "MAX_RULES_PER_RUN", 2)        # the Manual STIG holds 2, the checklist 5
    result = parse_stage([FIX / "evaluate_stig_results.xml"],
                         [FIX / "manual_stig_server2022.xml", FIX / "evaluate_stig_checklist.cklb"], tmp_path)
    limit = [w for w in result.warnings if "limit of 2 STIG reference rules" in w]
    assert len(limit) == 1 and limit[0].startswith("evaluate_stig_checklist.cklb: ")
    assert not any("0 rules" in w for w in result.warnings), result.warnings
    assert [s.file_name for s in result.enrichment.standalone_sources] == ["manual_stig_server2022.xml"]



def test_a_reference_with_a_commented_out_test_result_still_fills_text(tmp_path):
    text = (FIX / "manual_stig_win11.xml").read_text(encoding="utf-8")
    at = text.index("<Group")
    manual = tmp_path / "manual.xml"
    manual.write_text(text[:at] + "<!-- <TestResult> --><![CDATA[<TestResult>]]>" + text[at:], encoding="utf-8")
    result = parse_stage([FIX / "scc_embedded_results.xml"], [manual], tmp_path / "x")
    assert {f.vuln_id for f in result.findings if f.check_text} == {"V-253284", "V-253285"}
    assert not any("0 rule results" in w for w in result.warnings)


@pytest.mark.parametrize("as_reference", [False, True])
def test_the_advice_for_a_file_with_no_rule_results_knows_its_slot(tmp_path, as_reference):
    empty = tmp_path / "empty_results.xml"
    empty.write_text("<Benchmark id='b'><TestResult id='r'><target>H</target></TestResult></Benchmark>",
                     encoding="utf-8")
    results, references = ([FIX / "scc_embedded_results.xml"], [empty]) if as_reference else (
        [FIX / "scc_embedded_results.xml", empty], [])
    result = parse_stage(results, references, tmp_path / "x")
    line, = [w for w in result.warnings if w.startswith("empty_results.xml")]
    if as_reference:
        assert line == ("empty_results.xml: supplied as a STIG reference, but it holds a TestResult, so it was "
                        "read as scan results — 0 rule results, not counted as a scan")
    else:
        assert line == ("empty_results.xml: 0 rule results — not counted as a scan; if this file is a "
                        "benchmark, pass it as a reference")
