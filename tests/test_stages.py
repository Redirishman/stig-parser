from pathlib import Path

import pytest

from app.core.artifact_store import LocalArtifactStore
from app.core.findings_io import findings_from_json
from app.core.job_store import MemoryJobStore
from app.core.stages import (
    FINDINGS_KEY,
    INPUT_PREFIX,
    REPORT_KEY,
    run_export_stage,
    run_parse_stage,
)


def _seed_input(store, job_id, filename, data: bytes):
    store.put_bytes(f"{INPUT_PREFIX.format(job_id=job_id)}/{filename}", data)


def test_run_parse_stage_errors_on_garbage_input(tmp_path):
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    job_id = "job1"
    jobs.create(job_id, status="queued")
    _seed_input(store, job_id, "bad.xml", b"<html></html>")

    result = run_parse_stage(job_id, ["bad.xml"], store, jobs, work_dir=tmp_path / "w")
    assert result is False
    assert jobs.get(job_id)["status"] == "error"
    assert jobs.get(job_id)["error"]


def test_run_export_stage_reads_findings_and_writes_report(tmp_path):
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    job_id = "job2"
    jobs.create(job_id, status="running")

    # Seed a findings.json directly (bypassing parse) to test export in isolation.
    findings_json = (
        '[{"stig_title":"T","vuln_id":"V-1","rule_id":"r","severity":"CAT II",'
        '"status":"Open","server":"h","ip_address":"1.1.1.1","check_text":"c",'
        '"fix_text":"f"}]'
    )
    store.put_bytes(FINDINGS_KEY.format(job_id=job_id), findings_json.encode())

    ok = run_export_stage(job_id, store, jobs, work_dir=tmp_path / "w")
    assert ok is True
    assert store.exists(REPORT_KEY.format(job_id=job_id))
    job = jobs.get(job_id)
    assert job["status"] == "complete"
    assert job["summary"]["findings"] == 1


def test_run_parse_stage_rejects_traversal_filename(tmp_path):
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    job_id = "jobT"
    jobs.create(job_id, status="queued")

    outside = tmp_path / "outside_target.xml"
    result = run_parse_stage(
        job_id,
        ["../../outside_target.xml"],
        store,
        jobs,
        work_dir=tmp_path / "w",
    )
    assert result is False
    assert jobs.get(job_id)["status"] == "error"
    assert jobs.get(job_id)["error"] == "Invalid input filename."
    # Nothing was written outside the job work dir.
    assert not outside.exists()


def test_run_parse_stage_error_message_is_generic_on_unexpected_failure(tmp_path):
    # findings path is fine, but force an unexpected failure by pointing the
    # export stage at un-decodable findings JSON; assert no internal detail leaks.
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    job_id = "jobG"
    jobs.create(job_id, status="running")
    store.put_bytes(FINDINGS_KEY.format(job_id=job_id), b"not valid json{")

    ok = run_export_stage(job_id, store, jobs, work_dir=tmp_path / "w")
    assert ok is False
    err = jobs.get(job_id)["error"]
    assert err == "Export failed — see server logs."
    # Must not contain exception/library internals.
    assert "json" not in err.lower()
    assert "Traceback" not in err


def test_findings_key_roundtrips_through_store(tmp_path):
    store = LocalArtifactStore(tmp_path / "artifacts")
    job_id = "job3"
    findings_json = "[]"
    store.put_bytes(FINDINGS_KEY.format(job_id=job_id), findings_json.encode())
    loaded = findings_from_json(
        store.get_bytes(FINDINGS_KEY.format(job_id=job_id)).decode()
    )
    assert loaded == []


def test_run_parse_stage_reads_inputs_from_a_separate_store(tmp_path):
    """GovCloud keeps raw uploads in their own bucket (shorter retention), so
    inputs may come from a different store than the one findings are written to."""
    uploads = LocalArtifactStore(tmp_path / "uploads")
    artifacts = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    job_id = "job1"
    jobs.create(job_id, status="queued")

    scan = Path(__file__).parent / "fixtures" / "scc_results.xml"
    _seed_input(uploads, job_id, "scan.xml", scan.read_bytes())

    result = run_parse_stage(
        job_id,
        ["scan.xml"],
        artifacts,
        jobs,
        work_dir=tmp_path / "w",
        input_store=uploads,
    )

    assert result is True
    # Findings land in the artifacts store; the uploads store stays input-only.
    assert artifacts.exists(FINDINGS_KEY.format(job_id=job_id))
    assert not uploads.exists(FINDINGS_KEY.format(job_id=job_id))


def test_cancelled_job_does_not_enter_parse_stage(tmp_path):
    class ReadForbiddenStore(LocalArtifactStore):
        def size(self, key):
            raise AssertionError(f"cancelled job read input {key}")

    store = ReadForbiddenStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobC", status="cancelled", progress="Cancelled.")
    work_dir = tmp_path / "work"

    assert (
        run_parse_stage("jobC", ["scan.xml"], store, jobs, work_dir=work_dir) is False
    )
    assert jobs.get("jobC")["status"] == "cancelled"
    assert not work_dir.exists()


def test_export_completion_does_not_overwrite_midflight_cancel(tmp_path):
    jobs = MemoryJobStore()
    job_id = "jobC"

    class CancelOnUploadStore(LocalArtifactStore):
        def upload_from(self, key, source):
            super().upload_from(key, source)
            assert jobs.transition(job_id, "cancelled", progress="Cancelled.")

    store = CancelOnUploadStore(tmp_path / "artifacts")
    jobs.create(job_id, status="running", source_file_count=1)
    findings_json = (
        '[{"stig_title":"T","vuln_id":"V-1","rule_id":"r",'
        '"severity":"CAT II","status":"Open","server":"h",'
        '"ip_address":"1.1.1.1","check_text":"c","fix_text":"f"}]'
    )
    store.put_bytes(FINDINGS_KEY.format(job_id=job_id), findings_json.encode())

    assert run_export_stage(job_id, store, jobs, work_dir=tmp_path / "w") is False
    record = jobs.get(job_id)
    assert record["status"] == "cancelled"
    assert record["progress"] == "Cancelled."
    assert "summary" not in record


def test_duplicate_export_after_completion_is_idempotent(tmp_path):
    class ReadForbiddenStore(LocalArtifactStore):
        def get_bytes(self, key):
            raise AssertionError(f"completed job reread artifact {key}")

    store = ReadForbiddenStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobD", status="complete", summary={"findings": 1})

    assert run_export_stage("jobD", store, jobs, work_dir=tmp_path / "w") is True
    assert jobs.get("jobD")["status"] == "complete"


def test_parse_retry_after_parsed_phase_does_not_touch_inputs(tmp_path):
    class ReadForbiddenStore(LocalArtifactStore):
        def size(self, key):
            raise AssertionError(f"parsed job reread input {key}")

    jobs = MemoryJobStore()
    jobs.create("jobP", status="running", phase="parsed", progress="Parsed.")

    assert (
        run_parse_stage(
            "jobP",
            ["scan.xml"],
            ReadForbiddenStore(tmp_path / "artifacts"),
            jobs,
            work_dir=tmp_path / "w",
        )
        is True
    )
    assert jobs.get("jobP")["phase"] == "parsed"


def test_delayed_parse_failure_cannot_reclassify_exporting_job(tmp_path):
    jobs = MemoryJobStore()
    jobs.create("jobP", status="running", phase="parsing")

    class AdvanceThenFailStore(LocalArtifactStore):
        def size(self, key):
            assert jobs.update_if_status(
                "jobP",
                {"running"},
                expected_fields={"phase": "parsing"},
                phase="exporting",
            )
            raise RuntimeError("stale parser failed")

    assert (
        run_parse_stage(
            "jobP",
            ["scan.xml"],
            AdvanceThenFailStore(tmp_path / "artifacts"),
            jobs,
            work_dir=tmp_path / "w",
        )
        is True
    )
    assert jobs.get("jobP")["status"] == "running"
    assert jobs.get("jobP")["phase"] == "exporting"


def test_export_guard_losing_to_completion_is_idempotent(tmp_path):
    wrapped = MemoryJobStore()
    wrapped.create("jobR", status="running")

    class CompleteBeforeGuard:
        def __init__(self):
            self._injected = False

        def get(self, job_id):
            return wrapped.get(job_id)

        def update_if_status(self, job_id, expected_statuses, **fields):
            if not self._injected:
                self._injected = True
                assert wrapped.transition(
                    job_id, "complete", progress="Done.", summary={"findings": 1}
                )
            return wrapped.update_if_status(job_id, expected_statuses, **fields)

    class ReadForbiddenStore(LocalArtifactStore):
        def get_bytes(self, key):
            raise AssertionError(f"completed peer should prevent artifact read {key}")

    assert (
        run_export_stage(
            "jobR",
            ReadForbiddenStore(tmp_path / "artifacts"),
            CompleteBeforeGuard(),
            work_dir=tmp_path / "w",
        )
        is True
    )
    assert wrapped.get("jobR")["status"] == "complete"


# --- a parse stage error keeps the lines that explain it; the work directory never outlives the stage ---

def _empty_zip() -> bytes:
    import io
    import zipfile
    buf = io.BytesIO()
    zipfile.ZipFile(buf, "w").close()
    return buf.getvalue()


def test_a_parse_stage_error_record_carries_the_lines_that_explain_it(tmp_path):
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobW", status="queued")
    _seed_input(store, "jobW", "old.ckl.xml", b'<?xml version="1.0"?><CHECKLIST><ASSET/></CHECKLIST>')
    _seed_input(store, "jobW", "empty.zip", _empty_zip())
    assert run_parse_stage("jobW", ["old.ckl.xml", "empty.zip"], store, jobs, work_dir=tmp_path / "w") is False
    record = jobs.get("jobW")
    assert (record["status"], record["error"]) == ("error", "No valid results files could be parsed.")
    assert record["warnings"] == ['old.ckl.xml: STIG Viewer .ckl checklists are not supported — save it as .cklb in STIG Viewer 3 and upload that', 'No scan results or STIG references found in empty.zip (expected XCCDF results, a benchmark *xccdf.xml or *_Benchmark.xml, .cklb or .nessus files inside the zip)']


def test_a_parse_stage_error_record_keeps_at_most_200_clipped_lines(tmp_path, monkeypatch):
    import app.core.stages as stages
    from app.core.pipeline import PipelineError

    def failing_parse_stage(*args, **kwargs):
        raise PipelineError("No valid results files could be parsed.", [f"{i:03d}" + "x" * 5000 for i in range(250)])
    monkeypatch.setattr(stages, "parse_stage", failing_parse_stage)
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobB", status="queued")
    _seed_input(store, "jobB", "scan.xml", b"<a/>")
    assert run_parse_stage("jobB", ["scan.xml"], store, jobs, work_dir=tmp_path / "w") is False
    lines = jobs.get("jobB")["warnings"]
    assert len(lines) == 200
    assert lines[:199] == [f"{i:03d}" + "x" * 497 for i in range(199)]      # 500 characters each
    assert lines[199] == "… and 51 more warnings not shown"


def _scan_input(store, job_id):
    _seed_input(store, job_id, "scan.xml", (Path(__file__).parent / "fixtures" / "scc_results.xml").read_bytes())
    return ["scan.xml"]


def _bad_input(store, job_id):
    _seed_input(store, job_id, "bad.xml", b"<html></html>")
    return ["bad.xml"]


def _unsafe_name(store, job_id):
    return ["../outside.xml"]


@pytest.mark.parametrize("inputs, ok", [(_scan_input, True), (_bad_input, False), (_unsafe_name, False)])
def test_the_parse_stage_removes_its_work_directory_on_every_exit(tmp_path, inputs, ok):
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobR", status="queued")
    work_dir = tmp_path / "w"
    work_dir.mkdir()                                    # the Lambda makes it before the stage runs
    names = inputs(store, "jobR")
    assert run_parse_stage("jobR", names, store, jobs, work_dir=work_dir) is ok
    assert not work_dir.exists()


def test_the_parse_stage_removes_its_work_directory_after_an_unexpected_failure(tmp_path, monkeypatch):
    import app.core.stages as stages

    def boom(paths, *args, **kwargs):
        assert all(p.exists() for p in paths)           # the inputs were downloaded into the work dir
        raise RuntimeError("internal")
    monkeypatch.setattr(stages, "parse_stage", boom)
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobU", status="queued")
    names = _scan_input(store, "jobU")
    assert run_parse_stage("jobU", names, store, jobs, work_dir=tmp_path / "w") is False
    assert jobs.get("jobU")["error"] == "Parsing failed — see server logs."
    assert not (tmp_path / "w").exists()


# --- the reference hint splits the inputs; enrichment.json feeds the exporter ---------------------------

def test_parse_stage_uses_the_reference_hint_and_writes_enrichment_json(tmp_path):
    import json

    from openpyxl import load_workbook

    from app.core.stages import ENRICHMENT_KEY
    fix = Path(__file__).parent / "fixtures"
    store = LocalArtifactStore(tmp_path / "store")
    jobs = MemoryJobStore()
    job_id = "job-ref"
    jobs.create(job_id, status="queued")
    for name in ("evaluate_stig_results.xml", "evaluate_stig_checklist.cklb"):
        _seed_input(store, job_id, name, (fix / name).read_bytes())

    assert run_parse_stage(
        job_id, ["evaluate_stig_results.xml", "evaluate_stig_checklist.cklb"], store, jobs,
        work_dir=tmp_path / "w", reference_filenames=["evaluate_stig_checklist.cklb"],
    )
    assert jobs.get(job_id)["source_file_count"] == 1           # the checklist was a reference
    report = json.loads(store.get_bytes(ENRICHMENT_KEY.format(job_id=job_id)))
    assert any(s["file_name"] == "evaluate_stig_checklist.cklb" and not s["embedded"] for s in report["sources"])

    assert run_export_stage(job_id, store, jobs, work_dir=tmp_path / "w2")
    out = tmp_path / "r.xlsx"
    store.download_to(REPORT_KEY.format(job_id=job_id), out)
    assert any(row[0].value == "Reference sources" for row in load_workbook(out)["Summary"].iter_rows())


def test_without_the_hint_a_checklist_is_a_scan(tmp_path):
    fix = Path(__file__).parent / "fixtures"
    store = LocalArtifactStore(tmp_path / "store")
    jobs = MemoryJobStore()
    jobs.create("job-scan", status="queued")
    names = ["evaluate_stig_results.xml", "evaluate_stig_checklist.cklb"]
    for name in names:
        _seed_input(store, "job-scan", name, (fix / name).read_bytes())
    assert run_parse_stage("job-scan", names, store, jobs, work_dir=tmp_path / "w")
    assert jobs.get("job-scan")["source_file_count"] == 2




def test_a_parsed_job_record_keeps_at_most_200_clipped_warnings(tmp_path, monkeypatch):
    # One job record is one DynamoDB item (400 KB at most): a run's warnings are bounded on
    # success as they are on failure.
    import app.core.stages as stages
    from app.core.pipeline import ParseResult

    def parse_with_many_warnings(*args, **kwargs):
        return ParseResult(findings=[], warnings=[f"{i:03d}" + "x" * 5000 for i in range(250)],
                           source_file_count=1, coverage=set())
    monkeypatch.setattr(stages, "parse_stage", parse_with_many_warnings)
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobM", status="queued")
    names = _scan_input(store, "jobM")
    assert run_parse_stage("jobM", names, store, jobs, work_dir=tmp_path / "w") is True
    lines = jobs.get("jobM")["warnings"]
    assert len(lines) == 200
    assert lines[:199] == [f"{i:03d}" + "x" * 497 for i in range(199)]      # 500 characters each
    assert lines[199] == "… and 51 more warnings not shown"


def _seed_findings(store, job_id):
    store.put_bytes(FINDINGS_KEY.format(job_id=job_id), (
        '[{"stig_title":"T","vuln_id":"V-1","rule_id":"r","severity":"CAT II","status":"Open",'
        '"server":"h","ip_address":"1.1.1.1","check_text":"c","fix_text":"f"}]'
    ).encode())


def test_the_export_stage_leaves_no_workbook_in_its_work_directory(tmp_path):
    # A warm Lambda keeps /tmp between invocations; the workbook is in the store once uploaded.
    store = LocalArtifactStore(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobX", status="running")
    _seed_findings(store, "jobX")
    work_dir = tmp_path / "w"
    assert run_export_stage("jobX", store, jobs, work_dir=work_dir) is True
    assert store.exists(REPORT_KEY.format(job_id="jobX"))
    assert not list(work_dir.glob("*.xlsx"))


def test_the_export_stage_leaves_no_workbook_when_the_upload_fails(tmp_path):
    class UploadFails(LocalArtifactStore):
        def upload_from(self, key, path):
            raise OSError("store unavailable")

    store = UploadFails(tmp_path / "artifacts")
    jobs = MemoryJobStore()
    jobs.create("jobY", status="running")
    _seed_findings(store, "jobY")
    work_dir = tmp_path / "w"
    assert run_export_stage("jobY", store, jobs, work_dir=work_dir) is False
    assert jobs.get("jobY")["error"] == "Export failed — see server logs."
    assert not list(work_dir.glob("*.xlsx"))



def test_a_job_record_keeps_its_warnings_within_a_byte_budget():
    # One job record is one DynamoDB item (400 KB). The store writes the record as JSON, which escapes
    # non-ASCII text: an emoji costs 12 bytes there, so 200 lines of 500 emoji would be 1.2 MB.
    import json

    from app.core.stages import _MAX_RECORD_WARNING_BYTES, _bounded_warnings
    from app.reference.normalize import clip
    lines = [f"{i:03d} " + "\U0001F600" * 600 for i in range(200)]
    kept = _bounded_warnings(lines)
    assert len(json.dumps(kept)) <= _MAX_RECORD_WARNING_BYTES
    shown = kept[:-1]
    assert 0 < len(shown) < 200
    assert shown == [clip(line, 500) for line in lines[:len(shown)]]
    assert kept[-1] == f"… and {200 - len(shown)} more warnings not shown"


def test_plain_warnings_are_still_kept_200_at_a_time():
    from app.core.stages import _bounded_warnings
    lines = ["x" * 500] * 200
    assert _bounded_warnings(lines) == lines


class _NoListBucketStore(LocalArtifactStore):
    """S3 as the exporter's role sees it: s3:GetObject and s3:PutObject on jobs/*, no
    s3:ListBucket, so asking whether a missing key exists is refused (403), not answered."""

    def exists(self, key):
        if super().exists(key):
            return True
        from botocore.exceptions import ClientError
        raise ClientError({"Error": {"Code": "403", "Message": "Forbidden"}}, "HeadObject")


def _seed_parsed_job(store, jobs, job_id, **fields):
    """A job as a parse stage left it; *fields* are what that stage recorded on it."""
    jobs.create(job_id, status="running", phase="parsed", source_file_count=1, **fields)
    _seed_findings(store, job_id)


def test_a_job_parsed_before_enrichment_existed_exports_without_the_table(tmp_path):
    # A job parsed by the previous build has no enrichment.json and no flag: the exporter must
    # not ask S3 about the missing key, which its role cannot list (403, not 404).
    from openpyxl import load_workbook
    store = _NoListBucketStore(tmp_path / "store")
    jobs = MemoryJobStore()
    _seed_parsed_job(store, jobs, "old")
    assert run_export_stage("old", store, jobs, work_dir=tmp_path / "w") is True
    assert jobs.get("old")["status"] == "complete"
    out = tmp_path / "r.xlsx"
    store.download_to(REPORT_KEY.format(job_id="old"), out)
    assert not any(row[0].value == "Reference sources" for row in load_workbook(out)["Summary"].iter_rows())


def test_the_parse_stage_records_that_it_stored_enrichment_json(tmp_path):
    fix = Path(__file__).parent / "fixtures"
    store = LocalArtifactStore(tmp_path / "store")
    jobs = MemoryJobStore()
    jobs.create("new", status="queued")
    _seed_input(store, "new", "scc_embedded_results.xml", (fix / "scc_embedded_results.xml").read_bytes())
    assert run_parse_stage("new", ["scc_embedded_results.xml"], store, jobs, work_dir=tmp_path / "w")
    assert jobs.get("new")["enrichment_stored"] is True


def test_a_flagged_job_reads_enrichment_json_without_asking_whether_it_exists(tmp_path):
    import json

    from openpyxl import load_workbook

    from app.core.stages import ENRICHMENT_KEY
    from app.reference.enrich import EnrichmentReport
    from app.reference.models import ReferenceSource

    class NeverAsked(LocalArtifactStore):
        def exists(self, key):
            raise AssertionError(f"exists() asked for {key}")

    store = NeverAsked(tmp_path / "store")
    jobs = MemoryJobStore()
    _seed_parsed_job(store, jobs, "flagged", enrichment_stored=True)
    report = EnrichmentReport(sources=[ReferenceSource(file_name="manual.xml", benchmark_id="B", title="T",
                                                       release="V1R1", embedded=False, edition="manual",
                                                       rule_count=1)])
    store.put_bytes(ENRICHMENT_KEY.format(job_id="flagged"), json.dumps(report.to_dict()).encode())
    assert run_export_stage("flagged", store, jobs, work_dir=tmp_path / "w") is True
    out = tmp_path / "r.xlsx"
    store.download_to(REPORT_KEY.format(job_id="flagged"), out)
    files = [row[0].value for row in load_workbook(out)["Summary"].iter_rows()]
    assert "Reference sources" in files and "manual.xml" in files


def test_the_summary_line_always_fits_the_byte_budget():
    # Lines that fill the budget exactly, then one more: without the room kept for the
    # "… and N more" line, that line would push the record past the budget.
    import json

    from app.core.stages import _MAX_RECORD_WARNING_BYTES, _SUMMARY_RESERVE, _bounded_warnings, _stored_size
    full = "😀" * 500                                   # 6,004 stored bytes
    filler = "😀" * 155 + "abcd"                        # 1,868: 33 full lines + this = 200,000
    lines = [full] * 33 + [filler, "one more"]
    assert sum(_stored_size(line) for line in lines[:34]) == _MAX_RECORD_WARNING_BYTES
    kept = _bounded_warnings(lines)
    assert kept[-1] == "… and 2 more warnings not shown"
    assert len(json.dumps(kept)) <= _MAX_RECORD_WARNING_BYTES
    assert _SUMMARY_RESERVE >= _stored_size("… and 100000000 more warnings not shown")


def test_the_lambda_workbook_carries_the_warnings_of_the_parse(tmp_path):
    from tests.test_excel_exporter import run_warning_rows
    fix = Path(__file__).parent / "fixtures"
    store = LocalArtifactStore(tmp_path / "store")
    jobs = MemoryJobStore()
    jobs.create("warned", status="queued")
    _seed_input(store, "warned", "scc_embedded_results.xml", (fix / "scc_embedded_results.xml").read_bytes())
    _seed_input(store, "warned", "broken.xml", b"<TestResult><a></TestResult>")
    assert run_parse_stage("warned", ["scc_embedded_results.xml", "broken.xml"], store, jobs, work_dir=tmp_path / "w")
    assert run_export_stage("warned", store, jobs, work_dir=tmp_path / "w2")
    out = tmp_path / "r.xlsx"
    store.download_to(REPORT_KEY.format(job_id="warned"), out)
    rows = run_warning_rows(out)
    assert rows == jobs.get("warned")["warnings"]
    assert any(r.startswith("Could not parse results file: broken.xml") for r in rows)


def test_an_object_that_grew_after_its_size_was_checked_is_refused(tmp_path, monkeypatch):
    # The presigned PUT stays valid after the size check: an object can be replaced in between.
    import app.core.stages as stages

    class Replaced(LocalArtifactStore):
        def size(self, key):
            return 10                        # what the check saw; the bytes read are more

    fix = Path(__file__).parent / "fixtures"
    monkeypatch.setattr(stages, "MAX_UPLOAD_BYTES", 1000)
    store = Replaced(tmp_path / "store")
    jobs = MemoryJobStore()
    jobs.create("grown", status="queued")
    _seed_input(store, "grown", "scc_embedded_results.xml", (fix / "scc_embedded_results.xml").read_bytes())
    assert run_parse_stage("grown", ["scc_embedded_results.xml"], store, jobs, work_dir=tmp_path / "w") is False
    job = jobs.get("grown")
    assert job["status"] == "error"
    assert job["error"] == "File too large: 'scc_embedded_results.xml' (max 200 MB each)."
