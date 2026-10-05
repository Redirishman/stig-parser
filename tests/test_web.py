"""Tests for app.web — Flask routes, job pipeline, optional benchmarks."""
from __future__ import annotations

import io
import time
from pathlib import Path

import pytest

from app.web import create_app

FIXTURES = Path(__file__).parent / "fixtures"


def _make_upload(file_path: Path) -> tuple[io.BytesIO, str]:
    """Return (BytesIO, filename) suitable for Flask test-client multipart upload."""
    return io.BytesIO(file_path.read_bytes()), file_path.name


def _wait_for_completion(client, job_id: str, timeout: float = 10.0) -> dict:
    """Poll /api/status until the job leaves the running state or *timeout* elapses."""
    deadline = time.time() + timeout
    while time.time() < deadline:
        r = client.get(f"/api/status/{job_id}")
        data = r.get_json()
        if data.get("status") in ("complete", "error"):
            return data
        time.sleep(0.05)
    pytest.fail(f"Job {job_id} did not finish within {timeout}s")


@pytest.fixture()
def app(tmp_path, monkeypatch):
    """Build a Flask app whose job temp dir lives inside the test's tmp_path."""
    monkeypatch.setenv("STIG_TEMP_DIR", str(tmp_path / "jobs"))
    # Re-import so the new env var takes effect at module load.
    import importlib
    import app.web as web_module
    importlib.reload(web_module)
    flask_app = web_module.create_app(secret_key="test")
    flask_app.config["TESTING"] = True
    return flask_app


@pytest.fixture()
def client(app):
    return app.test_client()


# ---------------------------------------------------------------------------
# Static routes
# ---------------------------------------------------------------------------

class TestIndexRoute:
    def test_index_renders_200(self, client):
        r = client.get("/")
        assert r.status_code == 200
        assert b"STIG Compliance Parser" in r.data

    def test_index_marks_the_reference_zone_optional(self, client):
        r = client.get("/")
        # References are optional for every format; they add text the scan lacks.
        assert b'STIG References <span class="badge-optional">Optional</span>' in r.data


# ---------------------------------------------------------------------------
# /api/process validation
# ---------------------------------------------------------------------------

class TestProcessValidation:
    def test_no_results_returns_400(self, client):
        r = client.post("/api/process", data={})
        assert r.status_code == 400
        assert "No results" in r.get_json()["error"]

    def test_empty_results_filename_returns_400(self, client):
        data = {"results": (io.BytesIO(b""), "")}
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 400


# ---------------------------------------------------------------------------
# End-to-end job execution
# ---------------------------------------------------------------------------

class TestProcessWithSeparateBenchmark:
    def test_traditional_upload_completes(self, client):
        data = {
            "results": _make_upload(FIXTURES / "scc_results.xml"),
            "benchmarks": _make_upload(FIXTURES / "sample_benchmark.xml"),
        }
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 200
        job_id = r.get_json()["job_id"]

        final = _wait_for_completion(client, job_id)
        assert final["status"] == "complete", f"Job failed: {final}"

    def test_download_returns_xlsx(self, client):
        data = {
            "results": _make_upload(FIXTURES / "scc_results.xml"),
            "benchmarks": _make_upload(FIXTURES / "sample_benchmark.xml"),
        }
        post = client.post("/api/process", data=data, content_type="multipart/form-data")
        job_id = post.get_json()["job_id"]
        _wait_for_completion(client, job_id)

        r = client.get(f"/api/download/{job_id}")
        assert r.status_code == 200
        assert r.mimetype == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        # XLSX is a ZIP — magic bytes "PK"
        assert r.data[:2] == b"PK"


class TestProcessWithoutBenchmark:
    """SCC self-contained flow — no benchmark upload."""

    def test_results_only_upload_accepted(self, client):
        data = {"results": _make_upload(FIXTURES / "scc_results.xml")}
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 200
        assert "job_id" in r.get_json()

    def test_results_only_pipeline_completes(self, client):
        data = {"results": _make_upload(FIXTURES / "scc_results.xml")}
        post = client.post("/api/process", data=data, content_type="multipart/form-data")
        job_id = post.get_json()["job_id"]

        final = _wait_for_completion(client, job_id)
        # The fabricated SCC fixture has rule-results but no inline benchmark
        # definitions (it follows the 1.1 split-file pattern), so the pipeline
        # finishes — either complete with findings, or with an informative error
        # message about missing matches. Both are acceptable; what matters is
        # that no 400 was returned and the worker thread ran end to end.
        assert final["status"] in ("complete", "error")


# ---------------------------------------------------------------------------
# /api/status edge cases
# ---------------------------------------------------------------------------

class TestStatusRoute:
    def test_unknown_job_returns_404(self, client):
        r = client.get("/api/status/nonexistent-job-id")
        assert r.status_code == 404


class TestDownloadRoute:
    def test_unknown_job_returns_404(self, client):
        r = client.get("/api/download/nonexistent-job-id")
        assert r.status_code == 404


# ---------------------------------------------------------------------------
# /api/cancel
# ---------------------------------------------------------------------------

class TestCancelRoute:
    def test_unknown_job_returns_404(self, client):
        r = client.post("/api/cancel/nonexistent-job-id")
        assert r.status_code == 404

    def test_cancel_finished_job_reports_final_status(self, client):
        data = {"results": _make_upload(FIXTURES / "scc_results.xml")}
        post = client.post("/api/process", data=data, content_type="multipart/form-data")
        job_id = post.get_json()["job_id"]
        final = _wait_for_completion(client, job_id)

        r = client.post(f"/api/cancel/{job_id}")
        assert r.status_code == 200
        assert r.get_json()["status"] == final["status"]

    def test_cancel_running_job_sets_flag_and_worker_honors_it(self, client):
        import app.web as web

        job_id = "cancel-test-job"
        web._set_job(job_id, status="running", progress="working", warnings=[])
        with client.session_transaction() as sess:
            sess["job_id"] = job_id

        r = client.post(f"/api/cancel/{job_id}")
        assert r.status_code == 200
        assert r.get_json()["status"] == "cancelling"
        assert web._get_job(job_id).get("cancelled") is True

        with pytest.raises(web._JobCancelled):
            web._raise_if_cancelled(job_id)

    def test_status_endpoint_reports_cancelled(self, client):
        import app.web as web

        job_id = "cancelled-status-job"
        web._set_job(job_id, status="cancelled", progress="Cancelled.", warnings=[])
        with client.session_transaction() as sess:
            sess["job_id"] = job_id

        r = client.get(f"/api/status/{job_id}")
        assert r.get_json()["status"] == "cancelled"


# ---------------------------------------------------------------------------
# CKLB upload (Evaluate-STIG / STIG Viewer 3 checklists)
# ---------------------------------------------------------------------------

class TestCklbUpload:
    def test_cklb_only_upload_completes_with_findings(self, client):
        data = {"results": _make_upload(FIXTURES / "evaluate_stig_checklist.cklb")}
        post = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert post.status_code == 200
        job_id = post.get_json()["job_id"]

        final = _wait_for_completion(client, job_id)
        assert final["status"] == "complete"
        summary = final["summary"]
        assert summary["findings"] == 3
        assert summary["cat1"] == 1
        assert summary["cat2"] == 2  # includes the severity-override rule
        assert summary["hosts"] == 1

    def test_mixed_xml_and_cklb_upload_completes(self, client):
        data = {
            "results": [
                _make_upload(FIXTURES / "scc_results.xml"),
                _make_upload(FIXTURES / "evaluate_stig_checklist.cklb"),
            ]
        }
        post = client.post("/api/process", data=data, content_type="multipart/form-data")
        job_id = post.get_json()["job_id"]
        final = _wait_for_completion(client, job_id)
        assert final["status"] == "complete"
        assert final["summary"]["files"] == 2


class TestNessusUpload:
    def test_nessus_upload_completes_with_findings(self, client):
        data = {"results": _make_upload(FIXTURES / "nessus_compliance.nessus")}
        post = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert post.status_code == 200
        job_id = post.get_json()["job_id"]
        final = _wait_for_completion(client, job_id)
        assert final["status"] == "complete"
        summary = final["summary"]
        assert summary["findings"] == 4
        assert summary["cat1"] == 1
        assert summary["hosts"] == 1


# ---------------------------------------------------------------------------
# Upload allow-list + size cap (RESIDUALS #2)
# ---------------------------------------------------------------------------

class TestUploadValidation:
    def test_disallowed_extension_rejected(self, client):
        data = {"results": (io.BytesIO(b"malware"), "evil.exe")}
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 400
        assert "Unsupported file type" in r.get_json()["error"]

    def test_disallowed_benchmark_extension_rejected(self, client):
        data = {
            "results": _make_upload(FIXTURES / "scc_results.xml"),
            "benchmarks": (io.BytesIO(b"nope"), "notes.txt"),
        }
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 400
        assert "Unsupported file type" in r.get_json()["error"]

    def test_oversized_file_rejected(self, client, monkeypatch):
        import app.core.uploads as uploads
        monkeypatch.setattr(uploads, "MAX_UPLOAD_BYTES", 16)
        data = {"results": (io.BytesIO(b"x" * 64), "big.xml")}
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 400
        assert "too large" in r.get_json()["error"]

    def test_rejection_leaves_no_job_dir(self, client, tmp_path):
        data = {"results": (io.BytesIO(b"x"), "evil.exe")}
        client.post("/api/process", data=data, content_type="multipart/form-data")
        jobs_root = tmp_path / "jobs"
        # Validation runs before any mkdir — no orphan job directory created.
        assert not jobs_root.exists() or not any(jobs_root.iterdir())


# ---------------------------------------------------------------------------
# Security response headers (RESIDUALS #3)
# ---------------------------------------------------------------------------

class TestSecurityHeaders:
    def test_headers_present_on_index(self, client):
        r = client.get("/")
        assert r.headers["X-Frame-Options"] == "DENY"
        assert r.headers["X-Content-Type-Options"] == "nosniff"
        assert r.headers["Referrer-Policy"] == "no-referrer"
        csp = r.headers["Content-Security-Policy"]
        assert "default-src 'self'" in csp
        assert "script-src 'self'" in csp
        assert "'unsafe-inline'" not in csp

    def test_no_inline_script_in_template(self, client):
        # Strict CSP forbids inline JS — the page must load app.js externally.
        r = client.get("/")
        assert b"/static/app.js" in r.data


# ---------------------------------------------------------------------------
# Temp-dir cleanup (orphan reaping)
# ---------------------------------------------------------------------------

class TestJobCleanup:
    def test_errored_job_purges_its_dir(self, client, tmp_path):
        # A file that parses as XML-ish but yields no findings -> job errors.
        data = {"results": (io.BytesIO(b"<broken"), "bad.xml")}
        r = client.post("/api/process", data=data, content_type="multipart/form-data")
        assert r.status_code == 200
        job_id = r.get_json()["job_id"]

        final = _wait_for_completion(client, job_id)
        assert final["status"] == "error"

        # Status entry survives (poll returned "error"), but on-disk files are gone.
        job_dir = tmp_path / "jobs" / job_id
        assert not job_dir.exists()

    def test_sweep_removes_old_dirs_keeps_fresh(self, tmp_path, monkeypatch):
        import importlib
        import os
        monkeypatch.setenv("STIG_TEMP_DIR", str(tmp_path / "jobs"))
        import app.web as web_module
        importlib.reload(web_module)

        temp_dir = tmp_path / "jobs"
        temp_dir.mkdir(parents=True)
        old = temp_dir / "old-job"
        fresh = temp_dir / "fresh-job"
        old.mkdir()
        fresh.mkdir()

        # Age the old dir past the orphan cutoff.
        stale = time.time() - (web_module._ORPHAN_MAX_AGE_HOURS + 1) * 3600
        os.utime(old, (stale, stale))

        web_module._sweep_orphaned_jobs()

        assert not old.exists()
        assert fresh.exists()

    def test_sweep_spares_stale_dir_of_running_job(self, tmp_path, monkeypatch):
        """A job still 'running' must not be reaped no matter how stale its dir.

        Guards against the sweeper deleting a live worker's inputs out from under
        it — e.g. a slow parse of a huge archive that hasn't touched disk in >8h.
        """
        import importlib
        import os
        monkeypatch.setenv("STIG_TEMP_DIR", str(tmp_path / "jobs"))
        import app.web as web_module
        importlib.reload(web_module)

        temp_dir = tmp_path / "jobs"
        temp_dir.mkdir(parents=True)
        running = temp_dir / "running-job"
        orphan = temp_dir / "orphan-job"
        running.mkdir()
        orphan.mkdir()

        # Both dirs are equally stale; only status distinguishes them.
        stale = time.time() - (web_module._ORPHAN_MAX_AGE_HOURS + 1) * 3600
        os.utime(running, (stale, stale))
        os.utime(orphan, (stale, stale))

        # The running job has a live, non-terminal in-memory status; the orphan
        # has no entry at all (as if left by a prior process).
        web_module._set_job("running-job", status="running")

        web_module._sweep_orphaned_jobs()

        assert running.exists(), "sweeper deleted a running job's working dir"
        assert not orphan.exists()

        # A terminal status is fair game once stale (e.g. a completed, never
        # downloaded report).
        web_module._set_job("running-job", status="complete")
        web_module._sweep_orphaned_jobs()
        assert not running.exists()

    def test_start_orphan_sweeper_runs_sweep(self, tmp_path, monkeypatch):
        import importlib
        import os
        monkeypatch.setenv("STIG_TEMP_DIR", str(tmp_path / "jobs"))
        import app.web as web_module
        importlib.reload(web_module)

        temp_dir = tmp_path / "jobs"
        temp_dir.mkdir(parents=True)
        old = temp_dir / "old-job"
        old.mkdir()
        stale = time.time() - (web_module._ORPHAN_MAX_AGE_HOURS + 1) * 3600
        os.utime(old, (stale, stale))

        # A tiny interval means the daemon sweeps almost immediately.
        web_module._start_orphan_sweeper(interval_seconds=0.05)

        deadline = time.time() + 5.0
        while old.exists() and time.time() < deadline:
            time.sleep(0.05)
        assert not old.exists()


def test_a_logged_exception_reaches_the_operator_without_its_traceback():
    # The worker's log.exception for an unexpected failure is captured as a warning the
    # browser shows; the traceback names files and directories on the server.
    import logging

    from app.web import _WarningCollector
    lines: list[str] = []
    handler = _WarningCollector(lines)
    logger = logging.getLogger("app.test_collector")
    logger.addHandler(handler)
    try:
        try:
            raise OSError(2, "No such file or directory", "/srv/jobs/abc/member_1.xml")
        except OSError:
            logger.exception("Job %s failed with unhandled exception", "abc")
    finally:
        logger.removeHandler(handler)
    assert lines == ["Job abc failed with unhandled exception"]



def test_a_job_collects_only_the_warnings_of_its_own_worker_thread():
    # Every job's collector hangs on the one "app" logger: job A must never show job B's warnings.
    import logging
    import threading

    from app.web import _WarningCollector
    logger = logging.getLogger("app.test_two_jobs")
    both_ready = threading.Barrier(2)
    lines: dict[str, list[str]] = {"A": [], "B": []}

    def job(name: str) -> None:
        handler = _WarningCollector(lines[name])          # made in the worker thread, as _run_job does
        logger.addHandler(handler)
        try:
            both_ready.wait()
            for i in range(50):
                logger.warning("job %s: user-%s-secret-host_%d.xml", name, name, i)
            both_ready.wait()
        finally:
            logger.removeHandler(handler)

    workers = [threading.Thread(target=job, args=(name,)) for name in ("A", "B")]
    for worker in workers:
        worker.start()
    for worker in workers:
        worker.join()
    assert lines["A"] == [f"job A: user-A-secret-host_{i}.xml" for i in range(50)]
    assert lines["B"] == [f"job B: user-B-secret-host_{i}.xml" for i in range(50)]


def test_a_collector_that_cannot_tell_whose_a_warning_is_keeps_none_of_them(monkeypatch):
    # With thread ids switched off every record says thread None: the collector cannot tell job A's
    # warnings from job B's, and must keep none rather than show one job's file names in another.
    import logging
    import threading

    from app.web import _WarningCollector
    monkeypatch.setattr(logging, "logThreads", False)
    lines: list[str] = []
    handler = _WarningCollector(lines)
    logger = logging.getLogger("app.test_no_thread_ids")
    logger.addHandler(handler)
    try:
        other = threading.Thread(target=logger.warning, args=("job B: user-B-secret-host.xml",))
        other.start()
        other.join()
        logger.warning("job A: user-A-host.xml")
    finally:
        logger.removeHandler(handler)
    assert lines == []


def test_a_line_the_cap_already_counted_is_not_counted_again():
    # "line 200" is logged after the cap and counted in the "… and N more" line; the failed
    # run's own lines repeat it, and add one new line.
    import logging

    from app.web import _WarningCollector
    lines: list[str] = []
    handler = _WarningCollector(lines)
    logger = logging.getLogger("app.test_counted_once")
    logger.addHandler(handler)
    try:
        for i in range(201):
            logger.warning("line %d", i)
    finally:
        logger.removeHandler(handler)
    handler.add_new(["line 200", "line 7", "line 201"])
    assert lines == [f"line {i}" for i in range(200)] + ["… and 2 more warnings not shown"]


def test_what_the_collector_remembers_of_dropped_lines_is_bounded():
    # A flood of distinct lines past the cap (a 90,000-entry archive gave 89,800) is counted;
    # only the first 5,000 are remembered for de-duplication.
    import logging

    from app.web import _MAX_DROPPED_REMEMBERED, _WarningCollector
    assert _MAX_DROPPED_REMEMBERED == 5000
    lines: list[str] = []
    handler = _WarningCollector(lines)
    logger = logging.getLogger("app.test_flood")
    logger.addHandler(handler)
    try:
        for i in range(200 + 6000):
            logger.warning("flood %d", i)
    finally:
        logger.removeHandler(handler)
    assert len(handler._dropped_hashes) == 5000
    assert lines[-1] == "… and 6000 more warnings not shown"
    # Remembered lines are still not counted twice; past the memory, a repeat may be (the count
    # can then over-state, never under-state).
    handler.add_new(["flood 200", "flood 6199"])
    assert lines[-1] == "… and 6001 more warnings not shown"


def test_a_job_collects_at_most_200_warnings_and_says_how_many_more():
    import logging

    from app.web import _WarningCollector
    lines: list[str] = []
    handler = _WarningCollector(lines)
    logger = logging.getLogger("app.test_many_warnings")
    logger.addHandler(handler)
    try:
        for i in range(250):
            logger.warning("warning %d", i)
    finally:
        logger.removeHandler(handler)
    assert lines == [f"warning {i}" for i in range(200)] + ["… and 50 more warnings not shown"]



# --- uploads with one name are kept apart; extracted files do not outlive parsing ---------------------

def test_two_uploads_with_one_name_are_both_read_and_named_apart(client):
    data = {"results": [(io.BytesIO((FIXTURES / "scc_embedded_results.xml").read_bytes()), "results.xml"),
                        (io.BytesIO((FIXTURES / "evaluate_stig_results.xml").read_bytes()), "results.xml")]}
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "complete", final
    assert (final["summary"]["files"], final["summary"]["hosts"]) == (2, 2)
    # The Evaluate-STIG copy carries no benchmark: its line names it by the name it was uploaded with.
    assert any(w.startswith("results.xml (2): no matching STIG benchmark") for w in final["warnings"])


def _zipped_scan(path: Path) -> Path:
    import zipfile
    with zipfile.ZipFile(path, "w") as zf:
        zf.write(FIXTURES / "scc_embedded_results.xml", "scan.xml")
    return path


def _run(app, monkeypatch, *, cancelled: bool = False, parse=None):
    """Run one job in this thread over a zipped scan; the job record and its extraction dir."""
    import app.web as web
    job_id = "job-extract"
    upload = _zipped_scan(web._job_dir(job_id).parent / "upload.zip")
    web._set_job(job_id, status="running", warnings=[], cancelled=cancelled)
    if parse is not None:
        monkeypatch.setattr(web, "parse_stage", parse)
    web._run_job(job_id, [upload], [])
    return web._get_job(job_id), web._job_dir(job_id) / "benchmarks_extracted"


def test_the_extraction_dir_is_removed_after_a_successful_parse(app, monkeypatch, tmp_path):
    (tmp_path / "jobs").mkdir(exist_ok=True)
    job, extracted = _run(app, monkeypatch)
    assert job["status"] == "complete" and not extracted.exists()


def test_the_extraction_dir_is_removed_when_parsing_fails(app, monkeypatch, tmp_path):
    import app.web as web
    (tmp_path / "jobs").mkdir(exist_ok=True)

    def no_findings(results, references, extract_dir, **kwargs):
        extract_dir.mkdir(parents=True)
        raise web.PipelineError("No valid results files could be parsed.")

    job, extracted = _run(app, monkeypatch, parse=no_findings)
    assert job["status"] == "error" and not extracted.exists()


def test_the_extraction_dir_is_removed_after_an_unexpected_error(app, monkeypatch, tmp_path):
    (tmp_path / "jobs").mkdir(exist_ok=True)

    def broken(results, references, extract_dir, **kwargs):
        extract_dir.mkdir(parents=True)
        raise RuntimeError("bug")

    job, extracted = _run(app, monkeypatch, parse=broken)
    assert job["status"] == "error" and not extracted.exists()


def test_the_extraction_dir_is_removed_when_the_job_is_cancelled(app, monkeypatch, tmp_path):
    (tmp_path / "jobs").mkdir(exist_ok=True)
    job, extracted = _run(app, monkeypatch, cancelled=True)
    assert job["status"] == "cancelled" and not extracted.exists()


# --- a failed job shows the lines that explain it ------------------------------------------------------

def _empty_zip() -> bytes:
    import zipfile
    buf = io.BytesIO()
    zipfile.ZipFile(buf, "w").close()
    return buf.getvalue()


def test_a_failed_job_shows_the_lines_that_explain_it(client):
    data = {"results": [(io.BytesIO(b'<?xml version="1.0"?><CHECKLIST><ASSET/></CHECKLIST>'), "old.ckl.xml"),
                        (io.BytesIO(_empty_zip()), "empty.zip")]}
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "error", final
    assert final["error"] == "No valid results files could be parsed."
    assert 'old.ckl.xml: STIG Viewer .ckl checklists are not supported — save it as .cklb in STIG Viewer 3 and upload that' in final["warnings"]
    assert 'No scan results or STIG references found in empty.zip (expected XCCDF results, a benchmark *xccdf.xml or *_Benchmark.xml, .cklb or .nessus files inside the zip)' in final["warnings"]


def test_a_failed_job_shows_a_line_once_and_at_most_200_lines(client, monkeypatch):
    import logging

    import app.web as web_module
    from app.core.pipeline import PipelineError

    def failing_parse_stage(*args, **kwargs):
        logging.getLogger("app.test_failed_job").warning("line 0")      # also among the error's own lines
        raise PipelineError("No valid results files could be parsed.", [f"line {i}" for i in range(250)])
    monkeypatch.setattr(web_module, "parse_stage", failing_parse_stage)
    r = client.post("/api/process", data={"results": [_make_upload(FIXTURES / "scc_embedded_results.xml")]},
                    content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "error", final
    assert final["warnings"] == [f"line {i}" for i in range(200)] + ["… and 50 more warnings not shown"]


# --- the STIG References zone; the workbook carries the reference table ---------------------------------

def test_reference_zone_copy_and_accept_list(client):
    html = client.get("/").get_data(as_text=True)
    assert "STIG References" in html
    assert "Manual STIG ZIPs add check text" in html
    assert 'accept=".xml,.zip,.cklb"' in html
    assert "Not needed when uploading SCC result files" not in html
    assert "Add the Manual STIG as a reference to include check text" in html


def test_both_zones_take_a_zip_and_say_it_is_read_as_a_folder(client):
    html = client.get("/").get_data(as_text=True)
    assert 'accept=".xml,.zip,.cklb,.nessus"' in html             # the results zone
    assert html.count("A ZIP is read as a folder.") == 2
    script = client.get("/static/app.js").get_data(as_text=True)
    assert "['.xml', '.zip', '.cklb', '.nessus']" in script         # what a drop onto each zone keeps
    assert "['.xml', '.zip', '.cklb']" in script


def test_the_file_count_says_it_counts_result_files_read(client):
    html = client.get("/").get_data(as_text=True)
    assert "<dt>Result files read</dt>" in html
    assert "<dt>Result files</dt>" not in html


def test_manual_stig_reference_fills_check_text_through_the_web_flow(client, tmp_path):
    from openpyxl import load_workbook
    data = {
        "results": _make_upload(FIXTURES / "scc_embedded_results.xml"),
        "benchmarks": _make_upload(FIXTURES / "manual_stig_win11.xml"),
    }
    job_id = client.post("/api/process", data=data, content_type="multipart/form-data").get_json()["job_id"]
    status = _wait_for_completion(client, job_id)
    assert status["status"] == "complete", status
    assert any("not found in any supplied reference" in w for w in status["warnings"])
    out = tmp_path / "r.xlsx"
    out.write_bytes(client.get(f"/api/download/{job_id}").data)
    wb = load_workbook(out)
    assert wb["Findings"]["H2"].value                                  # check text filled
    assert wb["Findings"]["K2"].value.startswith("Check: manual_stig_win11.xml")
    assert any(row[0].value == "Reference sources" for row in wb["Summary"].iter_rows())


def test_cklb_is_accepted_in_the_reference_zone(client):
    data = {
        "results": _make_upload(FIXTURES / "evaluate_stig_results.xml"),
        "benchmarks": _make_upload(FIXTURES / "evaluate_stig_checklist.cklb"),
    }
    resp = client.post("/api/process", data=data, content_type="multipart/form-data")
    assert resp.status_code == 200
    final = _wait_for_completion(client, resp.get_json()["job_id"])
    assert final["status"] == "complete", final
    assert final["summary"]["files"] == 1                       # the checklist was a reference, not a scan


def _zip_bytes(members: dict[str, bytes]) -> bytes:
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, payload in members.items():
            zf.writestr(name, payload)
    return buf.getvalue()


@pytest.mark.parametrize("field", ["results", "benchmarks"])
@pytest.mark.parametrize("name", ["отчёт.zip", "测试.ZIP"])
def test_a_zip_whose_name_is_not_ascii_is_still_read_as_a_zip(client, field, name):
    # secure_filename keeps ASCII only: "отчёт.zip" was saved as "zip", with no extension,
    # and the archive was then sniffed as XML and not read.
    data = {} if field == "results" else {"results": _make_upload(FIXTURES / "evaluate_stig_results.xml")}
    data[field] = (io.BytesIO(_zip_bytes({"scan.xml": (FIXTURES / "scc_embedded_results.xml").read_bytes()})), name)
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "complete", final
    # The scan inside the ZIP was read (a scan in the reference zone is routed to the results).
    assert final["summary"]["hosts"] == (1 if field == "results" else 2), final


# --- a file that was not read: the run's own warning says why, once -------------------------------------

def _xml_error(payload: bytes, tmp_path) -> str:
    from lxml import etree

    from app.parsers.benchmark_parser import _safe_xml_parse
    from app.reference.normalize import error_text
    probe = tmp_path / "probe.xml"
    probe.write_bytes(payload)
    with pytest.raises(etree.XMLSyntaxError) as exc:
        _safe_xml_parse(probe)
    return error_text(exc.value)


def _json_error(payload: bytes) -> str:
    import json

    from app.reference.normalize import error_text
    with pytest.raises(ValueError) as exc:
        json.loads(payload.decode("utf-8"))
    return error_text(exc.value)


def _bad_crc_zip() -> bytes:
    import zipfile
    marker = b"MARKER-" + bytes(range(256))
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:
        zf.writestr("damaged.xml", b"<TestResult><!--" + marker + b"--></TestResult>")
    data = bytearray(buf.getvalue())
    data[data.find(marker) + 100] ^= 0xFF
    return bytes(data)


def broken_uploads():
    """The reviewer's set: one broken file of each kind, in each zone, next to a good scan."""
    results = [("broken.xml", b"<a><b></a>"), ("broken.cklb", b"{not json"),
               ("broken.nessus", b"<NessusClientData_v2><Report>"), ("notes.nessus", b"<foo/>"),
               ("broken.zip", b"PK\x03\x04garbage"), ("crc.zip", _bad_crc_zip())]
    references = [("broken_ref.xml", b"<Benchmark><unclosed>"), ("broken_ref.cklb", b"[1,2")]
    return results, references


def expected_reasons(tmp_path) -> dict[str, str]:
    return {
        "broken.xml": "Could not parse results file: broken.xml — invalid XML: "
                      + _xml_error(b"<a><b></a>", tmp_path),
        "broken.cklb": "Could not parse results file: broken.cklb — not valid JSON: " + _json_error(b"{not json"),
        "broken.nessus": "Could not parse results file: broken.nessus — invalid XML: "
                         + _xml_error(b"<NessusClientData_v2><Report>", tmp_path),
        "notes.nessus": "Could not parse results file: notes.nessus — not a Nessus export "
                        "(root element is <foo>, expected <NessusClientData_v2>)",
        "broken.zip": "broken.zip: not a readable ZIP — not read",
        "crc.zip": "crc.zip: could not read damaged.xml: Bad CRC-32 for file 'damaged.xml' — skipped",
        "broken_ref.xml": "Could not parse benchmark: broken_ref.xml — invalid XML: "
                          + _xml_error(b"<Benchmark><unclosed>", tmp_path),
        "broken_ref.cklb": "Could not parse reference checklist: broken_ref.cklb — not valid JSON: "
                           + _json_error(b"[1,2"),
    }


def test_a_file_that_was_not_read_is_named_once_with_the_reason(client, tmp_path):
    results, references = broken_uploads()
    data = {"results": [_make_upload(FIXTURES / "scc_embedded_results.xml")]
            + [(io.BytesIO(payload), name) for name, payload in results],
            "benchmarks": [(io.BytesIO(payload), name) for name, payload in references]}
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "complete", final
    for name, line in expected_reasons(tmp_path).items():
        about = [w for w in final["warnings"] if name in w]
        assert about == [line], (name, final["warnings"])


def test_a_failed_run_says_why_its_only_file_was_not_read(client, tmp_path):
    r = client.post("/api/process", data={"results": [(io.BytesIO(b"<a><b></a>"), "broken.xml")]},
                    content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "error", final
    assert final["warnings"] == [expected_reasons(tmp_path)["broken.xml"]]



# --- uploads keep their own names in the report, whatever name they are saved under ------------------------

def _scan_zip(extra: dict[str, bytes] | None = None) -> bytes:
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.write(FIXTURES / "scc_embedded_results.xml", "s/WKSTN-01.xml")
        zf.write(FIXTURES / "evaluate_stig_results.xml", "s/WIN-SERVER-04.xml")
        for name, payload in (extra or {}).items():
            zf.writestr(name, payload)
    return buf.getvalue()


def test_non_ascii_upload_names_are_the_names_the_report_shows(client, tmp_path):
    from openpyxl import load_workbook
    scan = (FIXTURES / "scc_embedded_results.xml").read_bytes().replace(b"WKSTN-01", b"WKSTN-07")
    data = {
        "results": [(io.BytesIO(_scan_zip({"s/notes.xml": b"<notes/>"})), "отчёт_сканирования.zip"),
                    (io.BytesIO(scan), "скан.xml"),
                    (io.BytesIO(b"<a><b></a>"), "сломан.xml")],
        "benchmarks": [(io.BytesIO((FIXTURES / "manual_stig_win11.xml").read_bytes()), "руководство.xml")],
    }
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    job_id = r.get_json()["job_id"]
    final = _wait_for_completion(client, job_id)
    assert final["status"] == "complete", final
    warnings = final["warnings"]
    assert ("отчёт_сканирования.zip: 1 XML file(s) were not recognised as scan results or STIG references: "
            "notes.xml") in warnings
    assert any(w.startswith("Could not parse results file: сломан.xml — invalid XML: ") for w in warnings)
    import re
    assert not any(re.search(r"\b(upload|reference)\.(xml|zip|cklb|nessus)\b", w) for w in warnings), warnings

    out = tmp_path / "r.xlsx"
    out.write_bytes(client.get(f"/api/download/{job_id}").data)
    wb = load_workbook(out)
    sources = [row[10].value for row in wb["Findings"].iter_rows(min_row=2)]
    assert any(s.startswith("Check: руководство.xml V2R9") for s in sources), sources
    files = [row[0].value for row in wb["Summary"].iter_rows()]
    assert "руководство.xml" in files and "reference.xml" not in files


def test_two_uploads_with_one_non_ascii_name_are_named_apart(client):
    data = {"results": [_make_upload(FIXTURES / "scc_embedded_results.xml"),
                        (io.BytesIO(b"<a><b></a>"), "скан.xml"),
                        (io.BytesIO(b"<c><d></c>"), "скан.xml")]}
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "complete", final
    named = sorted(w.split(" — ")[0] for w in final["warnings"] if w.startswith("Could not parse results file:"))
    assert named == ["Could not parse results file: скан.xml", "Could not parse results file: скан.xml (2)"]


def test_an_upload_is_saved_under_a_safe_name_and_named_by_its_own(client, monkeypatch):
    import app.web as web
    seen = {}

    def recording_parse_stage(results, references, extract_dir, **kwargs):
        seen["paths"] = [p.name for p in (*results, *references)]
        seen["names"] = sorted(kwargs.get("display_names", {}).values())
        raise web.PipelineError("No valid results files could be parsed.")
    monkeypatch.setattr(web, "parse_stage", recording_parse_stage)
    data = {"results": [(io.BytesIO(b"<a/>"), "отчёт.xml"), (io.BytesIO(b"<a/>"), "a b.xml")],
            "benchmarks": [(io.BytesIO(b"<a/>"), "руководство.zip")]}
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    _wait_for_completion(client, r.get_json()["job_id"])
    assert seen["paths"] == ["upload.xml", "a_b.xml", "reference.zip"]          # on disk: secure_filename's
    assert seen["names"] == ["a b.xml", "отчёт.xml", "руководство.zip"]          # to the operator: their own


def test_a_display_name_is_normalised_cleaned_and_bounded():
    from app.reference.normalize import display_file_name
    assert display_file_name("résumé.xml") == "résumé.xml"          # NFC
    assert display_file_name("a‮b\x07c .xml") == "abc.xml"                         # controls removed
    long = "x" * 300 + ".xml"
    assert len(display_file_name(long)) == 255 and display_file_name(long).endswith("x.xml")


# --- every problem is shown once on the card, with its reason (tests/problem_cases.py) -----------------

from tests.problem_cases import GOOD as CASE_GOOD, problem_cases  # noqa: E402

_CASES = problem_cases()


@pytest.mark.parametrize("case", _CASES, ids=[case.id for case in _CASES])
def test_each_problem_is_shown_once_on_the_card(client, case, monkeypatch):
    import app.utils.zip_extract as zip_extract
    for name, value in case.patches.items():
        monkeypatch.setattr(zip_extract, name, value)
    data = {"results": [_make_upload(CASE_GOOD)]
            + [(io.BytesIO(payload), name) for zone, name, payload in case.files if zone == "results"]}
    references = [(io.BytesIO(payload), name) for zone, name, payload in case.files if zone == "references"]
    if references:
        data["benchmarks"] = references
    r = client.post("/api/process", data=data, content_type="multipart/form-data")
    final = _wait_for_completion(client, r.get_json()["job_id"])
    assert final["status"] == "complete", final
    assert [w for w in final["warnings"] if case.subject in w] == case.shown


def test_the_web_workbook_carries_the_runs_warnings(client, tmp_path):
    from tests.test_excel_exporter import run_warning_rows
    data = {"results": [_make_upload(FIXTURES / "scc_embedded_results.xml"),
                        (io.BytesIO(b"<TestResult><a></TestResult>"), "broken.xml")]}
    job_id = client.post("/api/process", data=data, content_type="multipart/form-data").get_json()["job_id"]
    status = _wait_for_completion(client, job_id)
    assert status["status"] == "complete", status
    out = tmp_path / "r.xlsx"
    out.write_bytes(client.get(f"/api/download/{job_id}").data)
    assert any(r.startswith("Could not parse results file: broken.xml — invalid XML") for r in run_warning_rows(out))


def test_the_warnings_of_a_successful_run_respect_the_cap(app, monkeypatch, tmp_path):
    from app.core.pipeline import ParseResult
    from app.parsers.base import Finding
    (tmp_path / "jobs").mkdir(exist_ok=True)
    lines = [f"file{n}.xml: 0 rule results — not counted as a scan" for n in range(300)]

    def many_warnings(results, references, extract_dir, **kwargs):
        extract_dir.mkdir(parents=True)
        finding = Finding("T", "V-1", "SV-1r1_rule", "CAT II", "Open", "h", "1.1.1.1", "c", "f")
        return ParseResult(findings=[finding], warnings=lines, source_file_count=1, coverage=set())

    job, _ = _run(app, monkeypatch, parse=many_warnings)
    assert job["status"] == "complete"
    assert job["warnings"] == lines[:200] + ["… and 100 more warnings not shown"]


class TestErrorPurgesBeforeStatus:
    """A poller that sees "error" must never find the job's files still on disk.

    The uploads are scan data: once a job has failed they are useless, and the
    client is told the job is over the moment the status flips. Purging after
    the status is set left a window (and, for unexpected errors, left the files
    until the hourly sweep).
    """

    @staticmethod
    def _spy_on_error_status(monkeypatch):
        import app.web as web

        seen: dict[str, bool] = {}
        real_set_job = web._set_job

        def spy(job_id, **fields):
            if fields.get("status") == "error":
                seen["dir_existed_when_error_was_reported"] = web._job_dir(job_id).exists()
            return real_set_job(job_id, **fields)

        monkeypatch.setattr(web, "_set_job", spy)
        return seen

    def test_pipeline_error_purges_files_before_reporting_error(self, client, monkeypatch):
        seen = self._spy_on_error_status(monkeypatch)
        data = {"results": (io.BytesIO(b"<broken"), "bad.xml")}
        job_id = client.post(
            "/api/process", data=data, content_type="multipart/form-data"
        ).get_json()["job_id"]

        assert _wait_for_completion(client, job_id)["status"] == "error"
        assert seen == {"dir_existed_when_error_was_reported": False}

    def test_unexpected_error_purges_files_before_reporting_error(
        self, client, tmp_path, monkeypatch
    ):
        import app.web as web

        seen = self._spy_on_error_status(monkeypatch)

        def boom(*_args, **_kwargs):
            raise RuntimeError("boom")

        monkeypatch.setattr(web, "parse_stage", boom)
        data = {"results": _make_upload(FIXTURES / "scc_results.xml")}
        job_id = client.post(
            "/api/process", data=data, content_type="multipart/form-data"
        ).get_json()["job_id"]

        final = _wait_for_completion(client, job_id)
        assert final["status"] == "error"
        assert "boom" not in final["error"]
        assert seen == {"dir_existed_when_error_was_reported": False}
        assert not (tmp_path / "jobs" / job_id).exists()
